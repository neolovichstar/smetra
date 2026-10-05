"""Portable documents and bounded tabular transfers; no internal costs in PDFs."""

import base64
import csv
import io
import json
import zipfile
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape

try:
    from backend.business import DomainError, choice, identity, packed, stamp, string
except ModuleNotFoundError:
    from business import DomainError, choice, identity, packed, stamp, string


@dataclass
class Download:
    data: bytes
    mime: str
    filename: str


TITLES = {
    "estimate": "Смета",
    "proposal": "Коммерческое предложение",
    "invoice": "Счёт",
    "act": "Акт выполненных работ",
    "contract": "Договор - шаблон",
    "reference": "Справка о поступлениях",
}


def pdf(document):
    try:
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Spacer,
            Table,
            TableStyle,
            KeepTogether,
        )
    except ImportError:
        raise DomainError(
            503, "Генератор PDF не установлен. Установите requirements.txt на сервере"
        ) from None
    font = Path(__file__).resolve().parents[1] / "apps/web/assets/fonts/manrope-400.ttf"
    if "Smetra" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("Smetra", str(font)))
    data = json.loads(document["snapshot"])
    currency = data.get("currency", "RUB")

    def money(value):
        return f"{Decimal(value) / 100:,.2f}".replace(",", " ") + " " + currency

    palette = {
        "Minimal": "#1d2530",
        "Classic": "#3c4046",
        "Business": "#234564",
        "Modern": "#324968",
    }
    accent = colors.HexColor(palette[document["template"]])
    body = ParagraphStyle(
        "Body",
        fontName="Smetra",
        fontSize=9,
        leading=15,
        textColor=colors.HexColor("#343d49"),
        spaceAfter=8,
        wordWrap="CJK",
    )
    small = ParagraphStyle("Small", parent=body, fontSize=8, leading=12)
    heading = ParagraphStyle(
        "Heading", parent=body, fontSize=26, leading=34, textColor=accent, spaceAfter=18
    )
    label = ParagraphStyle(
        "Label", parent=body, fontSize=10, leading=16, textColor=accent
    )

    def p(text, style=body):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)

    stream = io.BytesIO()
    doc = SimpleDocTemplate(
        stream,
        pagesize=(595.28, 841.89),
        leftMargin=44,
        rightMargin=44,
        topMargin=48,
        bottomMargin=50,
        title=f"{document['name']} № {document['number']}",
        author=data["company"],
    )
    story = [
        p("СМЕТРА / " + data["company"], small),
        Spacer(1, 22),
        p(f"{document['name']} № {document['number']}", heading),
        p(data["title"], label),
        p("Клиент: " + data["client"]),
        p(data.get("company_details", "")),
        Spacer(1, 12),
        p(data.get("description", "")),
    ]
    rows = [
        [p("Позиция", small), p("Кол-во", small), p("Цена", small), p("Сумма", small)]
    ]
    for item in data.get("items", []):
        if not item["included"]:
            continue
        rows.append(
            [
                p(
                    item["name"]
                    + ("\n" + item["description"] if item.get("description") else ""),
                    small,
                ),
                p(item["quantity"] + " " + item["unit"], small),
                p(money(item["unit_price"]), small),
                p(money(item["subtotal"]), small),
            ]
        )
    if len(rows) == 1:
        rows.append(
            [
                p(data["title"], small),
                p("1", small),
                p(money(data["amount_kopecks"]), small),
                p(money(data["amount_kopecks"]), small),
            ]
        )
    table = Table(
        rows, colWidths=[230, 67, 103, 107], repeatRows=1, hAlign="LEFT", splitInRow=1
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#edf0f4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#dce1e8")),
                ("LEFTPADDING", (0, 0), (-1, -1), 9),
                ("RIGHTPADDING", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
            ]
        )
    )
    story += [
        Spacer(1, 10),
        table,
        Spacer(1, 20),
        KeepTogether(
            [
                p("Итого: " + money(data["amount_kopecks"]), label),
                p("Срок: " + data["due_date"] if data.get("due_date") else ""),
            ]
        ),
    ]
    if document["kind"] == "reference":
        story += [
            p("Записано поступлений: " + money(data.get("paid", 0)), label),
            p(
                "Справка об учтённых исполнителем поступлениях. Не является кассовым чеком."
            ),
        ]
    if document["kind"] == "contract":
        story += [
            p("Шаблон для согласования сторонами", label),
            p("Исполнитель: " + data["company"]),
            p("Заказчик: " + data["client"]),
            p("Предмет и состав работ указаны выше. Условия сторон:"),
            p(
                data.get("terms", "")
                or "Условия необходимо заполнить перед подписанием."
            ),
        ]
    else:
        story += [p(data.get("terms", ""))]
    if document["kind"] == "act":
        story += [
            Spacer(1, 20),
            p("Исполнитель: __________________    Заказчик: __________________"),
        ]
    footer = data.get("document_footer", "")
    for key, value in {
        "client.name": data["client"],
        "estimate.total": money(data["amount_kopecks"]),
        "project.name": data["title"],
        "company.name": data["company"],
    }.items():
        footer = footer.replace("{{" + key + "}}", str(value))
    story += [Spacer(1, 15), p(footer, small)]

    def page(canvas, pdfdoc):
        canvas.saveState()
        canvas.setFont("Smetra", 8)
        canvas.setFillColor(colors.HexColor("#7b8797"))
        canvas.drawString(44, 28, "Сметра · " + document["template"])
        canvas.drawRightString(551, 28, str(pdfdoc.page))
        canvas.restoreState()

    doc.build(story, onFirstPage=page, onLaterPages=page)
    return stream.getvalue()


def documents(service, method, parts, query, data):
    con = service.con
    if method == "GET":
        if parts:
            row = service.get("documents", parts[0])
            if len(parts) == 2 and parts[1] == "pdf":
                return 200, Download(
                    pdf(row), "application/pdf", f"smetra-{row['number']}.pdf"
                )
            return 200, {
                "document": {
                    k: row[k]
                    for k in ("id", "name", "kind", "template", "number", "created_at")
                }
            }
        return 200, {
            "items": [
                dict(r)
                for r in con.execute(
                    "SELECT id,name,kind,template,number,created_at FROM documents WHERE workspace_id=? ORDER BY created_at DESC LIMIT 50 OFFSET ?",
                    (service.wid, service.page(query)),
                )
            ]
        }
    if method != "POST":
        raise DomainError(405, "Документы сохраняются как неизменяемые снимки")
    document = prepare_document(service, data)
    service.insert("documents", document)
    service.emit("document", document["id"], "Документ сформирован", document["name"])
    return 201, {"document": {k: v for k, v in document.items() if k != "snapshot"}}


def prepare_document(service, data):
    con = service.con
    row = service.get("quotes", string(data.get("quote_id", ""), "Смета", 50, True))
    kind = choice(data.get("kind", "estimate"), TITLES, "Вид документа")
    template = choice(
        data.get("template", "Minimal"),
        ("Minimal", "Classic", "Business", "Modern"),
        "Шаблон",
    )
    snapshot = service.quote_view(row, bool(row["published_version"]))
    workspace = con.execute(
        "SELECT * FROM workspaces WHERE id=?", (service.wid,)
    ).fetchone()
    settings = json.loads(workspace["settings"])
    snapshot.update(
        company=workspace["name"],
        company_details=settings.get("company_details", ""),
        document_footer=settings.get("document_footer", ""),
    )
    project = con.execute(
        "SELECT * FROM projects WHERE quote_id=? AND workspace_id=?",
        (row["id"], service.wid),
    ).fetchone()
    if kind == "act" and (not project or project["status"] != "completed"):
        raise DomainError(409, "Акт создаётся после завершения заказа")
    if kind == "reference":
        if not project:
            raise DomainError(409, "Для справки нужен заказ с поступлениями")
        snapshot["paid"] = con.execute(
            "SELECT coalesce(sum(amount_kopecks),0) FROM project_payments WHERE project_id=?",
            (project["id"],),
        ).fetchone()[0]
    snapshot.pop("internal_cost", None)
    snapshot.pop("profit", None)
    snapshot.pop("custom_fields", None)
    for item in snapshot.get("items", []):
        for key in ("cost_price", "internal_cost", "markup"):
            item.pop(key, None)
    number = con.execute(
        "SELECT coalesce(max(number),0)+1 FROM documents WHERE workspace_id=?",
        (service.wid,),
    ).fetchone()[0]
    document = dict(
        id=identity(),
        workspace_id=service.wid,
        quote_id=row["id"],
        project_id=project["id"] if project else None,
        name=TITLES[kind],
        kind=kind,
        template=template,
        snapshot=packed(snapshot),
        number=number,
        created_at=stamp(),
    )
    pdf(document)  # Fail before persisting if rendering/dependencies are unavailable.
    return document


TRANSFER_COLUMNS = {
    "clients": ["name", "type", "company", "email", "phone", "telegram", "notes"],
    "catalog": [
        "name",
        "description",
        "unit",
        "price",
        "cost_price",
        "tax",
        "category",
    ],
}


def transfer(service, method, parts, query, data):
    if not parts or parts[0] not in TRANSFER_COLUMNS:
        raise DomainError(404, "Этот тип импорта не поддерживается")
    kind = parts[0]
    columns = TRANSFER_COLUMNS[kind]
    if method == "GET":
        table = "clients" if kind == "clients" else "catalog_items"
        records = service.con.execute(
            f"SELECT {','.join(columns)} FROM {table} WHERE workspace_id=? ORDER BY name LIMIT 10000",
            (service.wid,),
        ).fetchall()

        def safe(value):
            return (
                "'" + value
                if isinstance(value, str)
                and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r"))
                else value
            )

        rows = [columns] + [[safe(r[k]) for k in columns] for r in records]
        if query.get("format", ["csv"])[0] == "xlsx":
            try:
                from openpyxl import Workbook
            except ImportError:
                raise DomainError(503, "Установите openpyxl на сервере") from None
            book = Workbook(write_only=True)
            sheet = book.create_sheet("Сметра")
            for row in rows:
                sheet.append(row)
            stream = io.BytesIO()
            book.save(stream)
            return 200, Download(
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                kind + ".xlsx",
            )
        stream = io.StringIO(newline="")
        csv.writer(stream).writerows(rows)
        return 200, Download(
            stream.getvalue().encode("utf-8-sig"),
            "text/csv; charset=utf-8",
            kind + ".csv",
        )
    if method != "POST":
        raise DomainError(405, "Метод не поддерживается")
    service.write_access(True)
    if data.get("format") == "xlsx":
        try:
            raw = base64.b64decode(data.get("file", ""), validate=True)
            if len(raw) > 45000:
                raise ValueError("size")
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if (
                    sum(f.file_size for f in archive.infolist()) > 5_000_000
                    or len(archive.infolist()) > 100
                ):
                    raise ValueError("archive size")
                if any(
                    "vbaProject" in f.filename or "externalLinks" in f.filename
                    for f in archive.infolist()
                ):
                    raise ValueError("active content")
            from openpyxl import load_workbook

            book = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=False, keep_links=False
            )
            try:
                sheet = book.active
                if sheet.max_row > 201 or sheet.max_column > 30:
                    raise ValueError("row limit")
                values = list(sheet.iter_rows(values_only=True))
            finally:
                book.close()
            rows = [dict(zip(values[0], r)) for r in values[1:]]
        except (ValueError, KeyError, zipfile.BadZipFile, IndexError):
            raise DomainError(
                400, "Некорректный XLSX: до 200 строк, без макросов и внешних ссылок"
            ) from None
        except ImportError:
            raise DomainError(503, "Установите openpyxl на сервере") from None
    else:
        text = string(data.get("csv", ""), "CSV", 50000, True).lstrip("\ufeff")
        rows = list(csv.DictReader(io.StringIO(text)))
    if not 1 <= len(rows) <= 200:
        raise DomainError(400, "Импортируйте от 1 до 200 строк")
    seen_catalog = set()
    for index, record in enumerate(rows, start=2):
        payload = {
            k: (record[k] if record[k] is not None else "")
            for k in columns
            if k in record
        }
        if not payload.get("name"):
            raise DomainError(400, f"Строка {index}: нет названия (name)")
        for key in ("price", "cost_price"):
            if key in payload:
                try:
                    payload[key] = int(payload[key] or 0)
                except (ValueError, TypeError):
                    raise DomainError(
                        400, f"Строка {index}: цена должна быть целым числом копеек"
                    ) from None
        if kind == "catalog" and data.get("allow_duplicates") is not True:
            pair = (str(payload["name"]).strip().casefold(), str(payload.get("unit") or "шт.").strip().casefold())
            exists = service.con.execute(
                "SELECT 1 FROM catalog_items WHERE workspace_id=? AND lower(trim(name))=lower(trim(?)) "
                "AND lower(trim(unit))=lower(trim(?)) LIMIT 1",
                (service.wid, payload["name"], payload.get("unit") or "шт."),
            ).fetchone()
            if pair in seen_catalog or exists:
                raise DomainError(409, f"Строка {index}: похожая расценка уже есть. Разрешите дубликаты явно")
            seen_catalog.add(pair)
        try:
            service.entity(kind, "POST", [], {}, payload)
        except DomainError as err:
            raise DomainError(err.status, f"Строка {index}: {err.message}") from None
    service.emit(kind, service.wid, "Данные импортированы", str(len(rows)))
    return 201, {"count": len(rows)}
