"""Bounded, non-executing OOXML reader, used only in the isolated parser."""

import io
import csv
import re
import zipfile
import xml.etree.ElementTree as ET

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIMES = (DOCX, XLSX)
CSV = "text/csv"
VIEWABLE = (*MIMES, CSV)
READABLE = ("text/plain", "application/pdf", *VIEWABLE)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def archive(content, mime):
    if mime not in MIMES or not 0 < len(content) <= 5_000_000:
        raise ValueError("office bounds")
    result = {}
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        entries = package.infolist()
        if len(entries) > 512 or sum(entry.file_size for entry in entries) > 20_000_000:
            raise ValueError("archive bounds")
        for entry in entries:
            name = entry.filename
            lower = name.lower()
            if (name in result or name.startswith("/") or "\\" in name or
                    ".." in name.split("/") or ":" in name or entry.flag_bits & 1 or
                    entry.file_size > 5_000_000 or
                    entry.file_size > max(entry.compress_size, 1) * 200 or
                    any(part in lower for part in
                        ("vbaproject", "vbadata", "activex/", "embeddings/", "externallinks/"))):
                raise ValueError("unsafe archive")
            raw = package.read(entry)
            if lower.endswith((".xml", ".rels")):
                if re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", raw, re.I) or b"\x00" in raw:
                    raise ValueError("unsafe XML encoding or declarations")
                root = ET.fromstring(raw)
                if sum(1 for _ in root.iter()) > 100_000:
                    raise ValueError("XML bounds")
                if lower.endswith(".rels"):
                    for relation in root:
                        if (relation.get("TargetMode", "").lower() == "external" and
                                not relation.get("Type", "").endswith("/hyperlink")):
                            raise ValueError("external content")
                result[name] = root
            else:
                result[name] = None
    main = "word/document.xml" if mime == DOCX else "xl/workbook.xml"
    expected = ("application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
                if mime == DOCX else
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml")
    types = result.get("[Content_Types].xml")
    if main not in result or result[main] is None or types is None:
        raise ValueError("office structure")
    if not any(node.get("PartName") == "/" + main and node.get("ContentType") == expected for node in types):
        raise ValueError("office type mismatch")
    for node in types:
        if any(token in node.get("ContentType", "").lower() for token in ("macroenabled", "vba", "oleobject", "activex")):
            raise ValueError("active office content")
    return result


def preview(content, mime):
    if mime == CSV:
        if not 0 < len(content) <= 5_000_000:
            raise ValueError("CSV bounds")
        text = content.decode("utf-8-sig")
        if "\0" in text:
            raise ValueError("CSV text")
        try:
            dialect = csv.Sniffer().sniff(text[:10000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows, truncated = [], False
        budget = 120_000
        for number, row in enumerate(csv.reader(io.StringIO(text), dialect=dialect, strict=True), 1):
            if number > 100 or budget <= 0:
                truncated = True
                break
            values = []
            truncated |= len(row) > 20
            for value in row[:20]:
                length = min(1000, budget)
                truncated |= len(value) > length
                values.append(value[:length])
                budget -= len(values[-1])
            rows.append({"number": number, "cells": values})
        return {"format": "csv", "sheets": [{"name": "CSV", "rows": rows}], "truncated": truncated,
                "note": "До 100 строк и 20 столбцов. Ячейки показаны как текст. Оригинал сохраняется без изменений."}
    package = archive(content, mime)
    remaining = 120_000
    truncated = False

    def bounded(value, limit=1000):
        nonlocal remaining, truncated
        value = str(value) if value is not None else ""
        size = min(limit, remaining)
        truncated |= len(value) > size
        result = value[:size]
        remaining -= len(result)
        return result

    if mime == DOCX:
        body = package["word/document.xml"].find(W + "body")
        if body is None:
            raise ValueError("document body")
        blocks = []
        for index, node in enumerate(body):
            if index >= 400 or remaining <= 0:
                truncated = True
                break
            if node.tag == W + "p":
                value = bounded("".join(part.text or "" for part in node.iter(W + "t")), 10000)
                style = node.find(W + "pPr/" + W + "pStyle")
                heading = style is not None and (style.get(W + "val", "").lower().startswith("heading") or style.get(W + "val", "") == "Title")
                if value:
                    blocks.append({"type": "heading" if heading else "paragraph", "text": value})
            elif node.tag == W + "tbl":
                source_rows = node.findall(W + "tr")
                rows = []
                truncated |= len(source_rows) > 100
                for row in source_rows[:100]:
                    cells = row.findall(W + "tc")
                    truncated |= len(cells) > 20
                    rows.append([bounded("\n".join("".join(part.text or "" for part in p.iter(W + "t"))
                        for p in cell.findall(W + "p"))) for cell in cells[:20]])
                blocks.append({"type": "table", "rows": rows})
        return {"format": "docx", "blocks": blocks, "truncated": truncated,
                "note": "Текст и таблицы документа. Изображения, колонтитулы и исходная вёрстка доступны в скачанном оригинале."}

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        sheets = []
        truncated |= len(book.worksheets) > 6
        for sheet in book.worksheets[:6]:
            # The stored dimensions are untrusted. Explicit bounds also handle
            # missing dimensions, giant sparse ranges and malicious row indexes.
            rows = []
            truncated |= bool((sheet.max_row or 0) > 100 or (sheet.max_column or 0) > 20)
            for number, cells in enumerate(sheet.iter_rows(min_row=1, max_row=100, min_col=1, max_col=20), 1):
                values = []
                for cell in cells:
                    value = cell.value
                    if cell.data_type == "f":
                        value = "Формула (не рассчитана): " + str(value)
                    elif hasattr(value, "isoformat"):
                        value = value.isoformat()
                    values.append(bounded(value))
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    rows.append({"number": number, "cells": values})
                if remaining <= 0:
                    truncated = True
                    break
            sheets.append({"name": sheet.title, "rows": rows})
            if remaining <= 0:
                break
        return {"format": "xlsx", "sheets": sheets, "truncated": truncated,
                "note": "До 6 листов, 100 строк и 20 столбцов на лист. Формулы показаны как текст и не вычисляются. Оригинал сохраняется без изменений."}
    finally:
        book.close()


def text_sections(result):
    if result["format"] in ("xlsx", "csv"):
        return [(number, "Лист: " + sheet["name"] + "\n" + "\n".join(
            "Строка " + str(row["number"]) + ": " + " | ".join(row["cells"])
            for row in sheet["rows"])) for number, sheet in enumerate(result["sheets"], 1)]
    text = "\n".join(block["text"] if "text" in block else "\n".join(
        " | ".join(row) for row in block["rows"]) for block in result["blocks"])
    return [(number // 10000 + 1, text[number:number + 10000]) for number in range(0, len(text), 10000)]
