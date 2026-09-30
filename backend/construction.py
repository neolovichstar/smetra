"""Optional construction workflow beside the existing simple estimate flow."""

import base64
import binascii
import io
import re
import zipfile
from decimal import Decimal, ROUND_HALF_UP

from backend.business import DomainError, choice, date_value, identity, integer, stamp, string
from backend.construction_math import KINDS, UNITS, decimal_value, evaluate, length_in_metres, money, priced, zone_values
from backend.documents import Download


ZONE_KINDS = ("building", "floor", "room", "zone", "outdoor")
MEASURE_KINDS = ("length", "width", "height", "area", "volume", "count", "custom")
SOURCES = ("manual", "camera", "ar", "import", "ai", "calculated")
BUILTIN = frozenset(("length", "width", "height", "area", "perimeter", "volume", "openings"))


def row_in_object(service, table, item_id, object_id):
    row = service.con.execute(
        f"SELECT * FROM {table} WHERE id=? AND object_id=? AND workspace_id=?",
        (item_id, object_id, service.wid),
    ).fetchone()
    if not row:
        raise DomainError(404, "Запись объекта не найдена")
    return row


def variables(service, obj_id, zone=None):
    values = zone_values(
        zone["length_m"] if zone else 0,
        zone["width_m"] if zone else 0,
        zone["height_m"] if zone else 0,
        zone["openings_m2"] if zone else 0,
    )
    for row in service.con.execute(
        "SELECT symbol,value,unit,zone_id FROM construction_measurements "
        "WHERE object_id=? AND workspace_id=?",
        (obj_id, service.wid),
    ):
        if row["zone_id"] and (not zone or row["zone_id"] != zone["id"]):
            continue
        value = decimal_value(row["value"], row["symbol"])
        if row["unit"] in ("мм", "см", "м", "м.п."):
            value = length_in_metres(value, row["unit"])
        values[row["symbol"]] = value
    return values


def quantity_view(row, actual="0"):
    result = dict(row)
    result["planned_total_kopecks"] = priced(row["quantity"], row["unit_price"], row["price_coefficient"], row["markup_percent"], row["discount_percent"])
    result["planned_cost_kopecks"] = money(row["quantity"], row["cost_price"])
    result["actual_quantity"] = str(actual)
    result["actual_total_kopecks"] = priced(actual, row["unit_price"], row["price_coefficient"], row["markup_percent"], row["discount_percent"])
    return result


def calculate_row(service, obj, row):
    zone = row_in_object(service, "construction_zones", row["zone_id"], obj["id"]) if row["zone_id"] else None
    symbols = variables(service, obj["id"], zone)
    symbols.update(
        rate=decimal_value(row["consumption_rate"], "Норма расхода"),
        waste=decimal_value(row["waste_percent"], "Запас"),
        coefficient=decimal_value(row["coefficient"], "Коэффициент"),
    )
    if row["parent_work_id"]:
        parent = row_in_object(service, "construction_quantities", row["parent_work_id"], obj["id"])
        symbols["work_qty"] = decimal_value(parent["quantity"], "Объём работы")
    quantity = (evaluate(row["formula"], symbols) * symbols["coefficient"]).quantize(
        Decimal("0.0001"), rounding=ROUND_HALF_UP,
    )
    if quantity <= 0:
        raise DomainError(400, "Расчётный объём должен быть больше нуля")
    priced(quantity, row["unit_price"], row["price_coefficient"], row["markup_percent"], row["discount_percent"])
    return str(quantity)


def recalculate(service, obj):
    rows = service.con.execute(
        "SELECT * FROM construction_quantities WHERE object_id=? AND workspace_id=? ORDER BY created_at,id",
        (obj["id"], service.wid),
    ).fetchall()
    for row in sorted(rows, key=lambda item: bool(item["parent_work_id"])):
        quantity = calculate_row(service, obj, row)
        service.update("construction_quantities", row["id"], {
            "quantity": quantity, "updated_at": stamp(),
        })


def invalidate_quote(service, obj):
    if obj["quote_id"]:
        service.update("construction_objects", obj["id"], {
            "quote_id": None, "updated_at": stamp(),
        })


def detail(service, obj):
    zones = [dict(row) for row in service.con.execute(
        "SELECT * FROM construction_zones WHERE object_id=? AND workspace_id=? ORDER BY created_at,id",
        (obj["id"], service.wid),
    )]
    measurements = [dict(row) for row in service.con.execute(
        "SELECT * FROM construction_measurements WHERE object_id=? AND workspace_id=? ORDER BY created_at,id",
        (obj["id"], service.wid),
    )]
    rows = service.con.execute(
        "SELECT * FROM construction_quantities WHERE object_id=? AND workspace_id=? ORDER BY created_at,id",
        (obj["id"], service.wid),
    ).fetchall()
    facts = service.con.execute(
        "SELECT quantity_id,quantity FROM construction_facts WHERE object_id=? AND workspace_id=?",
        (obj["id"], service.wid),
    ).fetchall()
    actual = {}
    for fact in facts:
        actual[fact["quantity_id"]] = actual.get(fact["quantity_id"], Decimal(0)) + decimal_value(fact["quantity"], "Факт")
    quantities = [quantity_view(row, actual.get(row["id"], Decimal(0))) for row in rows]
    defects = [dict(row) for row in service.con.execute(
        "SELECT * FROM construction_defects WHERE object_id=? AND workspace_id=? "
        "ORDER BY created_at DESC,id DESC LIMIT 100",
        (obj["id"], service.wid),
    )]
    logs = [dict(row) for row in service.con.execute(
        "SELECT * FROM construction_daily_logs WHERE object_id=? AND workspace_id=? "
        "ORDER BY work_date DESC,created_at DESC,id DESC LIMIT 100",
        (obj["id"], service.wid),
    )]
    photos = {}
    for item in service.con.execute(
        "SELECT log_id,file_id FROM construction_log_photos WHERE object_id=? AND workspace_id=? "
        "ORDER BY created_at,file_id",
        (obj["id"], service.wid),
    ):
        photos.setdefault(item["log_id"], []).append(item["file_id"])
    for item in logs:
        item["photo_file_ids"] = photos.get(item["id"], [])
    return {
        "object": dict(obj), "zones": zones, "measurements": measurements,
        "quantities": quantities, "defects": defects, "daily_logs": logs,
        "totals": {
            "planned_kopecks": sum(item["planned_total_kopecks"] for item in quantities),
            "actual_kopecks": sum(item["actual_total_kopecks"] for item in quantities),
            "cost_kopecks": sum(item["planned_cost_kopecks"] for item in quantities),
        },
    }


def add_zone(service, obj, data):
    parent = data.get("parent_id") or None
    if parent:
        row_in_object(service, "construction_zones", parent, obj["id"])
    unit = choice(data.get("dimension_unit", "м"), ("мм", "см", "м"), "Единица размера")
    dimensions = {
        key: str(length_in_metres(data.get(key, 0), unit))
        for key in ("length", "width", "height")
    }
    openings = str(decimal_value(data.get("openings_m2", 0), "Площадь проёмов"))
    values = dict(
        id=identity(), workspace_id=service.wid, object_id=obj["id"], parent_id=parent,
        name=string(data.get("name", ""), "Помещение", 120, True),
        kind=choice(data.get("kind", "room"), ZONE_KINDS, "Тип зоны"),
        length_m=dimensions["length"], width_m=dimensions["width"],
        height_m=dimensions["height"], openings_m2=openings,
        notes=string(data.get("notes", ""), "Заметки", 2000),
        created_at=stamp(), updated_at=stamp(),
    )
    service.insert("construction_zones", values)
    service.emit("construction_zone", values["id"], "Помещение добавлено", values["name"])
    return 201, {"zone": values, "calculated": {key: str(value) for key, value in zone_values(
        values["length_m"], values["width_m"], values["height_m"], openings,
    ).items()}}


def add_measurement(service, obj, data):
    zone_id = data.get("zone_id") or None
    if zone_id:
        row_in_object(service, "construction_zones", zone_id, obj["id"])
    symbol = string(data.get("symbol", ""), "Переменная", 32, True)
    if not re.fullmatch(r"[a-z][a-z0-9_]{1,31}", symbol) or symbol in BUILTIN:
        raise DomainError(400, "Переменная: латинские буквы, цифры и _, без системных имён")
    if service.con.execute(
        "SELECT 1 FROM construction_measurements WHERE object_id=? AND symbol=?",
        (obj["id"], symbol),
    ).fetchone():
        raise DomainError(409, "Такая переменная уже есть на объекте")
    unit = choice(data.get("unit", "м"), UNITS, "Единица")
    value = str(decimal_value(data.get("value"), "Замер"))
    values = dict(
        id=identity(), workspace_id=service.wid, object_id=obj["id"], zone_id=zone_id,
        kind=choice(data.get("kind", "custom"), MEASURE_KINDS, "Вид замера"),
        symbol=symbol, value=value, unit=unit,
        source=choice(data.get("source", "manual"), SOURCES, "Источник"),
        notes=string(data.get("notes", ""), "Заметки", 2000),
        created_by=service.user["id"], created_at=stamp(),
    )
    service.insert("construction_measurements", values)
    service.emit("construction_measurement", values["id"], "Замер добавлен", symbol)
    return 201, {"measurement": values}


def add_quantity(service, obj, data):
    count = service.con.execute(
        "SELECT count(*) FROM construction_quantities WHERE object_id=? AND workspace_id=?",
        (obj["id"], service.wid),
    ).fetchone()[0]
    if count >= 200:
        raise DomainError(409, "В одном объекте может быть не больше 200 позиций сметы")
    zone_id = data.get("zone_id") or None
    zone = row_in_object(service, "construction_zones", zone_id, obj["id"]) if zone_id else None
    catalog_id = data.get("catalog_id") or None
    catalog = service.get("catalog_items", catalog_id) if catalog_id else None
    kind = choice(data.get("kind", catalog["item_type"] if catalog else "work"), KINDS, "Тип позиции")
    title = string(data.get("title", catalog["name"] if catalog else ""), "Работа или материал", 200, True)
    unit = choice(data.get("unit", catalog["unit"] if catalog else "м²"), UNITS, "Единица")
    parent_id = data.get("parent_work_id") or None
    parent = row_in_object(service, "construction_quantities", parent_id, obj["id"]) if parent_id else None
    if parent and (kind != "material" or parent["kind"] != "work"):
        raise DomainError(400, "Расход материала привязывается к работе")
    rate = decimal_value(data.get("consumption_rate", catalog["consumption_rate"] if catalog else 0), "Норма расхода")
    waste = decimal_value(data.get("waste_percent", 0), "Запас материала", maximum=Decimal(100))
    coefficient = decimal_value(data.get("coefficient", 1), "Коэффициент", minimum=Decimal("0.001"), maximum=Decimal(100))
    formula = string(data.get("formula", "work_qty * rate * (1 + waste / 100)" if parent else "area"), "Формула", 160, True)
    symbols = variables(service, obj["id"], zone)
    symbols.update(rate=rate, waste=waste, coefficient=coefficient)
    if parent:
        symbols["work_qty"] = decimal_value(parent["quantity"], "Объём работы")
    qty = evaluate(formula, symbols) * coefficient
    if qty <= 0:
        raise DomainError(400, "Объём должен быть больше нуля")
    qty = qty.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
    price = integer(data.get("unit_price", catalog["price"] if catalog else 0), "Расценка", 0, 100_000_000_000)
    cost = integer(data.get("cost_price", catalog["cost_price"] if catalog else 0), "Себестоимость", 0, 100_000_000_000)
    price_coefficient = decimal_value(data.get("price_coefficient", 1), "Коэффициент цены", minimum=Decimal("0.001"), maximum=Decimal(100))
    markup = decimal_value(data.get("markup_percent", 0), "Наценка", maximum=Decimal(1000))
    discount = decimal_value(data.get("discount_percent", 0), "Скидка", maximum=Decimal(100))
    priced(qty, price, price_coefficient, markup, discount)
    values = dict(
        id=identity(), workspace_id=service.wid, object_id=obj["id"], zone_id=zone_id,
        parent_work_id=parent_id, catalog_id=catalog_id, kind=kind, title=title,
        formula=formula, quantity=str(qty), unit=unit, unit_price=price,
        cost_price=cost, consumption_rate=str(rate), waste_percent=str(waste),
        coefficient=str(coefficient), notes=string(data.get("notes", ""), "Заметки", 2000),
        price_coefficient=str(price_coefficient), markup_percent=str(markup),
        discount_percent=str(discount),
        coefficient_reason=string(data.get("coefficient_reason", ""), "Причина коэффициента", 500),
        created_at=stamp(), updated_at=stamp(),
    )
    service.insert("construction_quantities", values)
    invalidate_quote(service, obj)
    service.emit("construction_quantity", values["id"], "Объём рассчитан", title)
    return 201, {"quantity": quantity_view(values)}


def add_fact(service, obj, data):
    row = row_in_object(
        service, "construction_quantities",
        string(data.get("quantity_id", ""), "Позиция", 80, True), obj["id"],
    )
    quantity = str(decimal_value(data.get("quantity"), "Факт", minimum=Decimal("0.0001")))
    values = dict(
        id=identity(), workspace_id=service.wid, object_id=obj["id"], quantity_id=row["id"],
        quantity=quantity, note=string(data.get("note", ""), "Примечание", 1000),
        created_by=service.user["id"], created_at=stamp(),
    )
    service.insert("construction_facts", values)
    service.emit("construction_fact", values["id"], "Факт зафиксирован", row["title"])
    return 201, {"fact": values}


def defect_fields(service, obj, data, previous=None):
    zone_id = data.get("zone_id", previous["zone_id"] if previous else None) or None
    if zone_id:
        row_in_object(service, "construction_zones", zone_id, obj["id"])
    photo_id = data.get("photo_file_id", previous["photo_file_id"] if previous else None) or None
    if photo_id:
        photo = service.get("files", photo_id)
        if photo["construction_id"] != obj["id"] or photo["mime"] not in ("image/png", "image/jpeg"):
            raise DomainError(400, "Нужно фото именно этого объекта")
    return dict(
        zone_id=zone_id, photo_file_id=photo_id,
        description=string(data.get("description", previous["description"] if previous else ""), "Описание дефекта", 2000, True),
        severity=choice(data.get("severity", previous["severity"] if previous else "normal"), ("low", "normal", "high"), "Важность"),
        measurement_note=string(data.get("measurement_note", previous["measurement_note"] if previous else ""), "Замер дефекта", 500),
        suggested_work=string(data.get("suggested_work", previous["suggested_work"] if previous else ""), "Предлагаемая работа", 500),
        status=choice(data.get("status", previous["status"] if previous else "open"), ("open", "in_progress", "resolved"), "Статус дефекта"),
        updated_at=stamp(),
    )


def log_fields(service, obj, data, previous=None):
    work_date = date_value(data.get("work_date", previous["work_date"] if previous else ""), "Дата работ")
    if not work_date:
        raise DomainError(400, "Укажите дату работ")
    zone_id = data.get("zone_id", previous["zone_id"] if previous else None) or None
    if zone_id:
        row_in_object(service, "construction_zones", zone_id, obj["id"])
    quantity_id = data.get("quantity_id", previous["quantity_id"] if previous else None) or None
    quantity_row = row_in_object(service, "construction_quantities", quantity_id, obj["id"]) if quantity_id else None
    if quantity_row and not zone_id:
        zone_id = quantity_row["zone_id"]
    if quantity_row and zone_id and quantity_row["zone_id"] and quantity_row["zone_id"] != zone_id:
        raise DomainError(400, "Работа относится к другому помещению")
    completed = decimal_value(data.get("completed_quantity", previous["completed_quantity"] if previous else 0),
                              "Выполненный объём", maximum=Decimal("1000000"))
    unit = quantity_row["unit"] if quantity_row else choice(
        data.get("unit", previous["unit"] if previous else ""), UNITS | {""}, "Единица объёма",
    )
    if completed > 0 and not unit:
        raise DomainError(400, "Укажите единицу выполненного объёма")
    return dict(
        zone_id=zone_id, quantity_id=quantity_id, work_date=work_date,
        workers=string(data.get("workers", previous["workers"] if previous else ""), "Исполнители", 500),
        worker_count=integer(data.get("worker_count", previous["worker_count"] if previous else 0), "Количество людей", 0, 100),
        work_description=string(data.get("work_description", previous["work_description"] if previous else ""), "Выполненная работа", 2000, True),
        completed_quantity=str(completed), unit=unit,
        comment=string(data.get("comment", previous["comment"] if previous else ""), "Комментарий", 2000),
        updated_at=stamp(),
    )


def validate_log_photo(service, obj, file_id):
    photo = service.get("files", string(file_id, "Фото", 80, True))
    if photo["construction_id"] != obj["id"] or photo["mime"] not in ("image/png", "image/jpeg"):
        raise DomainError(400, "Фото должно принадлежать этому объекту")
    return photo["id"]


def add_log_fact(service, obj, fields):
    if not fields["quantity_id"] or decimal_value(fields["completed_quantity"], "Объём") == 0:
        return None
    _, result = add_fact(service, obj, {
        "quantity_id": fields["quantity_id"], "quantity": fields["completed_quantity"],
        "note": "Журнал работ: " + fields["work_date"],
    })
    return result["fact"]["id"]


def create_daily_log(service, obj, data):
    count = service.con.execute(
        "SELECT count(*) FROM construction_daily_logs WHERE object_id=? AND workspace_id=?",
        (obj["id"], service.wid),
    ).fetchone()[0]
    if count >= 5000:
        raise DomainError(409, "На объекте может быть не больше 5000 записей журнала")
    photos = data.get("photo_file_ids", [])
    if (not isinstance(photos, list) or len(photos) > 4
            or not all(isinstance(item, str) for item in photos)
            or len(set(photos)) != len(photos)):
        raise DomainError(400, "Укажите не больше четырёх разных фото")
    photos = [validate_log_photo(service, obj, file_id) for file_id in photos]
    fields = log_fields(service, obj, data)
    values = dict(id=identity(), workspace_id=service.wid, object_id=obj["id"],
                  fact_id=add_log_fact(service, obj, fields),
                  created_by=service.user["id"], created_at=stamp(), **fields)
    service.insert("construction_daily_logs", values)
    for file_id in photos:
        service.con.execute(
            "INSERT INTO construction_log_photos(workspace_id,object_id,log_id,file_id,created_at) VALUES(?,?,?,?,?)",
            (service.wid, obj["id"], values["id"], file_id, stamp()),
        )
    service.emit("construction_daily_log", values["id"], "Запись журнала добавлена", values["work_description"][:120])
    return 201, {"daily_log": dict(values, photo_file_ids=photos)}


def as_xlsx(service, obj):
    from openpyxl import Workbook

    snapshot = detail(service, obj)
    book = Workbook()
    sheet = book.active
    sheet.title = "Ведомость"
    sheet.append(["Вид", "Работа / материал", "Помещение", "Формула", "Объём", "Ед.", "Расценка, ₽", "План, ₽", "Факт", "Работа-основание", "Норма", "Запас, %", "Коэффициент объёма", "Коэффициент цены", "Наценка, %", "Скидка, %"])
    zone_names = {zone["id"]: zone["name"] for zone in snapshot["zones"]}
    quantity_names = {item["id"]: item["title"] for item in snapshot["quantities"]}
    def safe_cell(value):
        value = str(value or "")
        return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value
    for item in snapshot["quantities"]:
        sheet.append([
            item["kind"], safe_cell(item["title"]), safe_cell(zone_names.get(item["zone_id"], "")), safe_cell(item["formula"]),
            float(item["quantity"]), item["unit"], item["unit_price"] / 100,
            item["planned_total_kopecks"] / 100, float(item["actual_quantity"]),
            safe_cell(quantity_names.get(item["parent_work_id"], "")),
            float(item["consumption_rate"]), float(item["waste_percent"]), float(item["coefficient"]),
            float(item["price_coefficient"]), float(item["markup_percent"]), float(item["discount_percent"]),
        ])
    sheet.freeze_panes = "A2"
    for column, width in {"A": 16, "B": 38, "C": 24, "D": 34, "E": 16, "F": 10, "G": 18, "H": 18, "I": 14, "J": 30, "K": 14, "L": 14, "M": 18, "N": 18, "O": 14, "P": 14}.items():
        sheet.column_dimensions[column].width = width
    output = io.BytesIO()
    book.save(output)
    return 200, Download(output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "smetra-vedomost.xlsx")


def import_xlsx(service, obj, data):
    from openpyxl import load_workbook

    encoded = data.get("content", "")
    if not isinstance(encoded, str) or len(encoded) > 2_000_000:
        raise DomainError(413, "XLSX слишком большой")
    try:
        raw = base64.b64decode(encoded, validate=True)
        if not 1 <= len(raw) <= 1_500_000:
            raise ValueError("size")
        archive = zipfile.ZipFile(io.BytesIO(raw))
        members = archive.infolist()
        if len(members) > 100 or sum(item.file_size for item in members) > 5_000_000:
            raise ValueError("expanded size")
        if any(item.filename.startswith("/") or ".." in item.filename.split("/") for item in members):
            raise ValueError("path")
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        sheet = book.active
        if sheet.max_row > 201 or sheet.max_column > 20:
            raise ValueError("shape")
        rows = list(sheet.iter_rows(min_row=2, max_row=201, max_col=16, values_only=True))
        if any(isinstance(value, str) and value.startswith("=") for row in rows for value in row):
            raise ValueError("formulas")
    except (ValueError, TypeError, binascii.Error, zipfile.BadZipFile, KeyError, OSError, AttributeError):
        raise DomainError(400, "Некорректная ведомость XLSX") from None
    finally:
        if "book" in locals():
            book.close()
    zones = {row["name"]: row["id"] for row in service.con.execute(
        "SELECT id,name FROM construction_zones WHERE object_id=? AND workspace_id=?",
        (obj["id"], service.wid),
    )}
    rows = [row for row in rows if any(value is not None for value in row)]
    if not rows or len(rows) > 200:
        raise DomainError(400, "В ведомости должно быть от 1 до 200 позиций")
    existing = service.con.execute(
        "SELECT count(*) FROM construction_quantities WHERE object_id=? AND workspace_id=?",
        (obj["id"], service.wid),
    ).fetchone()[0]
    if existing + len(rows) > 200:
        raise DomainError(409, "В объекте может быть не больше 200 позиций")
    def restore_cell(value):
        text = str(value or "")
        return text[1:] if text.startswith("'") and text[1:].lstrip().startswith(("=", "+", "-", "@")) else text

    imported = {}
    for row in sorted(rows, key=lambda item: item[0] == "material"):
        kind, title, zone_name, formula, _, unit, price, _, _, parent_name, rate, waste, coefficient, price_coefficient, markup, discount = row
        title, zone_name, formula, parent_name = map(restore_cell, (title, zone_name, formula, parent_name))
        if zone_name and zone_name not in zones:
            raise DomainError(400, "Помещение из XLSX не найдено в объекте")
        if parent_name and parent_name not in imported:
            raise DomainError(400, "Работа-основание из XLSX не найдена")
        price_value = decimal_value(price if price is not None else 0, "Цена")
        cents = price_value * 100
        if cents != cents.to_integral_value():
            raise DomainError(400, "Цена XLSX должна быть указана с точностью до копейки")
        payload = {
            "kind": kind, "title": title, "zone_id": zones.get(zone_name),
            "parent_work_id": imported.get(parent_name),
            "formula": formula, "unit": unit,
            "unit_price": int(cents), "consumption_rate": str(rate or 0),
            "waste_percent": str(waste or 0), "coefficient": str(coefficient or 1),
            "price_coefficient": str(price_coefficient or 1),
            "markup_percent": str(markup or 0), "discount_percent": str(discount or 0),
        }
        _, result = add_quantity(service, obj, payload)
        if kind == "work":
            if title in imported:
                raise DomainError(400, "Названия работ-оснований должны быть уникальными")
            imported[title] = result["quantity"]["id"]
    service.emit("construction_object", obj["id"], "Ведомость XLSX импортирована", str(len(rows)))
    return 201, {"imported": len(rows)}


def to_quote(service, obj):
    if obj["quote_id"]:
        return 200, {"quote": service.quote_view(service.get("quotes", obj["quote_id"]))}
    snapshot = detail(service, obj)
    if not snapshot["quantities"]:
        raise DomainError(409, "Добавьте работы или материалы в ведомость")
    if len(snapshot["quantities"]) > 200:
        raise DomainError(409, "Для одной сметы выберите не более 200 позиций")
    client = service.get("clients", obj["client_id"]) if obj["client_id"] else None
    items = [{
        "name": item["title"], "quantity": item["quantity"], "unit": item["unit"],
        "unit_price": item["unit_price"], "cost_price": item["cost_price"],
        "coefficient": item["price_coefficient"], "markup": item["markup_percent"],
        "discount": item["discount_percent"],
        "category": item["kind"], "description": item["notes"],
    } for item in snapshot["quantities"]]
    quote = service.create_quote({
        "title": obj["name"][:120], "client": client["name"] if client else "Клиент не указан",
        "client_id": obj["client_id"], "description": obj["description"], "items": items,
    })
    service.con.execute(
        "UPDATE construction_objects SET quote_id=?,updated_at=? WHERE id=? AND workspace_id=?",
        (quote["id"], stamp(), obj["id"], service.wid),
    )
    service.emit("construction_object", obj["id"], "Смета создана из ведомости", quote["id"])
    return 201, {"quote": quote}


def route(service, method, parts, query, data):
    if parts == ["units"] and method == "GET":
        return 200, {"units": sorted(UNITS), "kinds": KINDS}
    if parts == ["calculate"] and method == "POST":
        symbols = zone_values(data.get("length", 0), data.get("width", 0), data.get("height", 0), data.get("openings", 0))
        return 200, {"quantity": str(evaluate(data.get("formula", "area"), symbols)),
                     "variables": {key: str(value) for key, value in symbols.items()}}
    if not parts or parts[0] != "objects":
        raise DomainError(404, "Объект не найден")
    if len(parts) == 1:
        if method == "GET":
            return 200, {"items": [dict(row) for row in service.con.execute(
                "SELECT id,name,description,status,client_id,quote_id,project_id,created_at,updated_at "
                "FROM construction_objects WHERE workspace_id=? ORDER BY updated_at DESC LIMIT 50 OFFSET ?",
                (service.wid, service.page(query)),
            )]}
        if method == "POST":
            client_id = data.get("client_id") or None
            if client_id:
                service.get("clients", client_id)
            values = dict(
                id=identity(), workspace_id=service.wid, client_id=client_id, quote_id=None,
                project_id=None, name=string(data.get("name", ""), "Объект", 200, True),
                description=string(data.get("description", ""), "Описание", 5000),
                status="survey", created_at=stamp(), updated_at=stamp(),
            )
            service.insert("construction_objects", values)
            service.emit("construction_object", values["id"], "Объект создан", values["name"])
            return 201, {"object": values}
        raise DomainError(405, "Метод не поддерживается")
    obj = service.get("construction_objects", parts[1])
    rest = parts[2:]
    if not rest:
        if method == "GET":
            return 200, detail(service, obj)
        if method == "PATCH":
            name = string(data.get("name", obj["name"]), "Объект", 200, True)
            description = string(data.get("description", obj["description"]), "Описание", 5000)
            service.update("construction_objects", obj["id"], dict(name=name,description=description,updated_at=stamp()))
            return 200, {"object": dict(obj, name=name, description=description)}
        raise DomainError(405, "Метод не поддерживается")
    if rest == ["zones"] and method == "POST":
        return add_zone(service, obj, data)
    if len(rest) == 2 and rest[0] == "zones" and method == "PATCH":
        zone = row_in_object(service, "construction_zones", rest[1], obj["id"])
        unit = choice(data.get("dimension_unit", "м"), ("мм", "см", "м"), "Единица размера")
        values = {
            "name": string(data.get("name", zone["name"]), "Помещение", 120, True),
            "length_m": str(length_in_metres(data["length"], unit)) if "length" in data else zone["length_m"],
            "width_m": str(length_in_metres(data["width"], unit)) if "width" in data else zone["width_m"],
            "height_m": str(length_in_metres(data["height"], unit)) if "height" in data else zone["height_m"],
            "openings_m2": str(decimal_value(data.get("openings_m2", zone["openings_m2"]), "Площадь проёмов")),
            "notes": string(data.get("notes", zone["notes"]), "Заметки", 2000),
            "updated_at": stamp(),
        }
        service.update("construction_zones", zone["id"], values)
        recalculate(service, obj)
        invalidate_quote(service, obj)
        service.emit("construction_zone", zone["id"], "Размеры помещения уточнены")
        return 200, detail(service, service.get("construction_objects", obj["id"]))
    if rest == ["measurements"] and method == "POST":
        return add_measurement(service, obj, data)
    if len(rest) == 2 and rest[0] == "measurements" and method == "PATCH":
        measure = row_in_object(service, "construction_measurements", rest[1], obj["id"])
        values = {
            "value": str(decimal_value(data.get("value", measure["value"]), "Замер")),
            "unit": choice(data.get("unit", measure["unit"]), UNITS, "Единица"),
            "notes": string(data.get("notes", measure["notes"]), "Заметки", 2000),
        }
        service.update("construction_measurements", measure["id"], values)
        recalculate(service, obj)
        invalidate_quote(service, obj)
        service.emit("construction_measurement", measure["id"], "Замер уточнён")
        return 200, detail(service, service.get("construction_objects", obj["id"]))
    if rest == ["quantities"] and method == "POST":
        return add_quantity(service, obj, data)
    if len(rest) == 2 and rest[0] == "quantities" and method == "PATCH":
        row = row_in_object(service, "construction_quantities", rest[1], obj["id"])
        values = {
            "title": string(data.get("title", row["title"]), "Позиция", 200, True),
            "formula": string(data.get("formula", row["formula"]), "Формула", 160, True),
            "unit": choice(data.get("unit", row["unit"]), UNITS, "Единица"),
            "unit_price": integer(data.get("unit_price", row["unit_price"]), "Цена", 0, 100_000_000_000),
            "cost_price": integer(data.get("cost_price", row["cost_price"]), "Себестоимость", 0, 100_000_000_000),
            "consumption_rate": str(decimal_value(data.get("consumption_rate", row["consumption_rate"]), "Норма расхода")),
            "waste_percent": str(decimal_value(data.get("waste_percent", row["waste_percent"]), "Запас", maximum=Decimal(100))),
            "coefficient": str(decimal_value(data.get("coefficient", row["coefficient"]), "Коэффициент", minimum=Decimal("0.001"), maximum=Decimal(100))),
            "price_coefficient": str(decimal_value(data.get("price_coefficient", row["price_coefficient"]), "Коэффициент цены", minimum=Decimal("0.001"), maximum=Decimal(100))),
            "markup_percent": str(decimal_value(data.get("markup_percent", row["markup_percent"]), "Наценка", maximum=Decimal(1000))),
            "discount_percent": str(decimal_value(data.get("discount_percent", row["discount_percent"]), "Скидка", maximum=Decimal(100))),
            "coefficient_reason": string(data.get("coefficient_reason", row["coefficient_reason"]), "Причина коэффициента", 500),
            "notes": string(data.get("notes", row["notes"]), "Заметки", 2000),
            "updated_at": stamp(),
        }
        service.update("construction_quantities", row["id"], values)
        recalculate(service, obj)
        invalidate_quote(service, obj)
        service.emit("construction_quantity", row["id"], "Позиция пересчитана")
        return 200, detail(service, service.get("construction_objects", obj["id"]))
    if rest == ["facts"] and method == "POST":
        return add_fact(service, obj, data)
    if rest == ["defects"] and method == "POST":
        count = service.con.execute(
            "SELECT count(*) FROM construction_defects WHERE object_id=? AND workspace_id=?",
            (obj["id"], service.wid),
        ).fetchone()[0]
        if count >= 100:
            raise DomainError(409, "На объекте может быть не больше 100 дефектов")
        values = dict(id=identity(), workspace_id=service.wid, object_id=obj["id"],
                      created_by=service.user["id"], created_at=stamp(),
                      **defect_fields(service, obj, data))
        service.insert("construction_defects", values)
        service.emit("construction_defect", values["id"], "Дефект добавлен", values["description"][:120])
        return 201, {"defect": values}
    if len(rest) == 2 and rest[0] == "defects" and method == "PATCH":
        defect = row_in_object(service, "construction_defects", rest[1], obj["id"])
        values = defect_fields(service, obj, data, defect)
        service.update("construction_defects", defect["id"], values)
        service.emit("construction_defect", defect["id"], "Дефект обновлён", values["status"])
        return 200, {"defect": dict(defect, **values)}
    if rest == ["logs"] and method == "POST":
        return create_daily_log(service, obj, data)
    if len(rest) >= 2 and rest[0] == "logs":
        log = row_in_object(service, "construction_daily_logs", rest[1], obj["id"])
        if len(rest) == 2 and method == "PATCH":
            fields = log_fields(service, obj, data, log)
            if log["fact_id"]:
                service.con.execute(
                    "DELETE FROM construction_facts WHERE id=? AND workspace_id=?",
                    (log["fact_id"], service.wid),
                )
            fields["fact_id"] = add_log_fact(service, obj, fields)
            service.update("construction_daily_logs", log["id"], fields)
            service.emit("construction_daily_log", log["id"], "Запись журнала изменена")
            return 200, {"daily_log": dict(log, **fields)}
        if len(rest) == 2 and method == "DELETE":
            service.con.execute(
                "DELETE FROM construction_daily_logs WHERE id=? AND workspace_id=?",
                (log["id"], service.wid),
            )
            if log["fact_id"]:
                service.con.execute(
                    "DELETE FROM construction_facts WHERE id=? AND workspace_id=?",
                    (log["fact_id"], service.wid),
                )
            service.emit("construction_daily_log", log["id"], "Ошибочная запись журнала удалена")
            return 200, {"ok": True}
        if len(rest) == 3 and rest[2] == "photos" and method == "POST":
            file_id = validate_log_photo(service, obj, data.get("file_id"))
            found = service.con.execute(
                "SELECT 1 FROM construction_log_photos WHERE log_id=? AND file_id=? AND workspace_id=?",
                (log["id"], file_id, service.wid),
            ).fetchone()
            if found:
                return 200, {"ok": True}
            count = service.con.execute(
                "SELECT count(*) FROM construction_log_photos WHERE log_id=? AND workspace_id=?",
                (log["id"], service.wid),
            ).fetchone()[0]
            if count >= 4:
                raise DomainError(409, "К записи можно прикрепить не больше четырёх фото")
            service.con.execute(
                "INSERT INTO construction_log_photos(workspace_id,object_id,log_id,file_id,created_at) VALUES(?,?,?,?,?)",
                (service.wid, obj["id"], log["id"], file_id, stamp()),
            )
            service.emit("construction_daily_log", log["id"], "Фото добавлено к записи журнала")
            return 201, {"ok": True}
        if len(rest) == 4 and rest[2] == "photos" and method == "DELETE":
            service.con.execute(
                "DELETE FROM construction_log_photos WHERE log_id=? AND file_id=? AND workspace_id=?",
                (log["id"], rest[3], service.wid),
            )
            service.emit("construction_daily_log", log["id"], "Фото удалено из записи журнала")
            return 200, {"ok": True}
    if len(rest) == 2 and rest[0] == "facts" and method == "DELETE":
        fact = row_in_object(service, "construction_facts", rest[1], obj["id"])
        if service.con.execute(
            "SELECT 1 FROM construction_daily_logs WHERE fact_id=? AND workspace_id=?",
            (fact["id"], service.wid),
        ).fetchone():
            raise DomainError(409, "Этот факт записан через журнал. Исправьте или удалите запись журнала")
        service.con.execute("DELETE FROM construction_facts WHERE id=? AND workspace_id=?", (fact["id"], service.wid))
        service.emit("construction_fact", fact["id"], "Ошибочная запись факта удалена")
        return 200, {"ok": True}
    if rest == ["quote"] and method == "POST":
        return to_quote(service, obj)
    if rest == ["xlsx"] and method == "GET":
        return as_xlsx(service, obj)
    if rest == ["import"] and method == "POST":
        return import_xlsx(service, obj, data)
    raise DomainError(404, "Действие не найдено")
