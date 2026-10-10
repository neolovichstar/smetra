import concurrent.futures
import base64
import contextlib
import datetime as dt
import importlib.util
import io
import json
import os
import tempfile
import threading
import unittest
import urllib.error
import urllib.parse
import urllib.request
from unittest.mock import patch
from http.server import ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace

from backend.business import DomainError, Service, calculate


class CalculationTests(unittest.TestCase):
    def test_receipt_ocr_asks_for_final_total_without_spending_tokens_on_reasoning(self):
        from backend.receipt_ocr import _request

        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def read(self, _size):
                return json.dumps({"choices": [{"message": {"content": (
                    '```json\n{"merchant":"Магазин","amount_kopecks":174800,'
                    '"date":"2026-10-01"}\n```'
                )}}]}).encode()

        with patch.dict(os.environ, OPENROUTER_API_KEY="test-only"), patch(
            "backend.receipt_ocr.urllib.request.urlopen", return_value=Response()
        ) as opener:
            result = _request("openrouter/free", "data:image/png;base64,abc")
        payload = json.loads(opener.call_args.args[0].data)
        self.assertEqual(result["amount_kopecks"], 174800)
        self.assertEqual(payload["reasoning"], {"effort": "none"})
        self.assertIn("grand total", payload["messages"][0]["content"])

    def test_moscow_month_boundary_uses_local_business_date(self):
        from backend.business import business_date, business_month_start

        moment = int(dt.datetime(2026, 9, 30, 23, 30, tzinfo=dt.timezone.utc).timestamp())
        self.assertEqual(business_date(moment), dt.date(2026, 10, 1))
        self.assertEqual(
            business_month_start(moment),
            int(dt.datetime(2026, 9, 30, 21, 0, tzinfo=dt.timezone.utc).timestamp()),
        )

    def test_construction_formula_is_bounded_and_safe(self):
        from backend.construction_math import evaluate, zone_values

        values = zone_values(5, 4, 3, 2)
        self.assertEqual(str(evaluate('area - openings', values)), '18.0000')
        for expression in ('__import__("os")', '2 ** 1000', '1 / 0', '-1', 'missing'):
            with self.subTest(expression=expression), self.assertRaises(DomainError):
                evaluate(expression, values)

    def test_construction_pricing_matches_quote_order(self):
        from backend.construction_math import priced

        self.assertEqual(priced('20', 15000, '1.2', '15', '10'), 372600)
        items, total, _ = calculate([{
            'name': 'Painting', 'quantity': '20', 'unit_price': 15000,
            'coefficient': '1.2', 'markup': '15', 'discount': '10',
        }])
        self.assertEqual(total, 372600)
        self.assertEqual(items[0]['coefficient'], '1.2')

    def test_decimal_rounding_discount_markup_tax_and_optional(self):
        items, total, cost = calculate(
            [
                {
                    "name": "Work",
                    "quantity": "1.5",
                    "unit_price": 10001,
                    "cost_price": 4000,
                    "markup": "10",
                    "discount": "5",
                    "tax": "20",
                },
                {
                    "name": "Extra",
                    "quantity": "1",
                    "unit_price": 99900,
                    "optional": True,
                    "included": False,
                },
            ]
        )
        self.assertEqual(items[0]["subtotal"], 18812)
        self.assertEqual(total, 18812)
        self.assertEqual(cost, 6000)

    def test_rejects_nonfinite_bool_overflow_and_empty(self):
        for value in ("NaN", "Infinity", "-1", "0", "0.00001", True):
            with self.subTest(value=value), self.assertRaises(DomainError):
                calculate([{"name": "x", "quantity": value, "unit_price": 100}])
        for items in (
            [],
            [{"name": "x", "unit_price": True}],
            [{"name": "x", "quantity": "1000000", "unit_price": 10000000000}],
        ):
            with self.assertRaises(DomainError):
                calculate(items)


class BusinessFlows(unittest.TestCase):
    def test_construction_bulk_price_preview_apply_undo_and_conflict(self):
        owner, _ = self.account('bulk-owner')
        outsider, _ = self.account('bulk-outsider')
        _, created = self.call('/construction/objects', 'POST', {'name': 'Bulk prices'}, owner)
        obj = created['object']['id']
        _, zone = self.call(f'/construction/objects/{obj}/zones', 'POST', {
            'name': 'Room', 'length': '5', 'width': '4',
        }, owner)
        zone_id = zone['zone']['id']
        ids = []
        for title, price in [('Painting', 10000), ('Plaster', 20000)]:
            code, result = self.call(f'/construction/objects/{obj}/quantities', 'POST', {
                'zone_id': zone_id, 'kind': 'work', 'title': title,
                'formula': 'area', 'unit_price': price,
            }, owner)
            self.assertEqual(code, 201, result)
            ids.append(result['quantity']['id'])
        code, old_quote = self.call(f'/construction/objects/{obj}/quote', 'POST', {}, owner)
        self.assertEqual(code, 201, old_quote)
        endpoint = f'/construction/objects/{obj}/prices'
        payload = {'ids': ids, 'percent': '8'}
        self.assertEqual(self.call(endpoint + '/preview', 'POST', payload, outsider)[0], 404)
        code, preview = self.call(endpoint + '/preview', 'POST', payload, owner)
        self.assertEqual(code, 200, preview)
        self.assertEqual(sorted(item['new_unit_price'] for item in preview['changes']), [10800, 21600])
        before = self.call(f'/construction/objects/{obj}', token=owner)[1]
        self.assertEqual(sorted(item['unit_price'] for item in before['quantities']), [10000, 20000])
        code, applied = self.call(endpoint + '/apply', 'POST', {
            **payload, 'expected_hash': preview['expected_hash'],
        }, owner)
        self.assertEqual(code, 200, applied)
        self.assertTrue(applied['quote_recreation_required'])
        self.assertEqual(self.call(endpoint + '/apply', 'POST', {
            **payload, 'expected_hash': preview['expected_hash'],
        }, owner)[0], 409)
        changed = self.call(f'/construction/objects/{obj}', token=owner)[1]
        self.assertEqual(sorted(item['unit_price'] for item in changed['quantities']), [10800, 21600])
        self.assertIsNone(changed['object']['quote_id'])
        self.assertEqual(self.call('/quotes/' + old_quote['quote']['id'], token=owner)[0], 200)
        batch_id = applied['batch_id']
        self.assertEqual(changed['latest_price_batch']['id'], batch_id)
        self.assertEqual(self.call(endpoint + '/batches/' + batch_id, 'POST', {}, outsider)[0], 404)
        self.assertEqual(self.call(endpoint + '/batches/' + batch_id, 'POST', {}, owner)[0], 200)
        self.assertEqual(self.call(endpoint + '/batches/' + batch_id, 'POST', {}, owner)[0], 409)
        restored = self.call(f'/construction/objects/{obj}', token=owner)[1]
        self.assertEqual(sorted(item['unit_price'] for item in restored['quantities']), [10000, 20000])
        code, preview = self.call(endpoint + '/preview', 'POST', payload, owner)
        self.assertEqual(code, 200, preview)
        code, applied = self.call(endpoint + '/apply', 'POST', {
            **payload, 'expected_hash': preview['expected_hash'],
        }, owner)
        self.assertEqual(code, 200, applied)
        self.assertEqual(self.call(f'/construction/objects/{obj}/quantities/{ids[0]}', 'PATCH', {
            'unit_price': 14000,
        }, owner)[0], 200)
        self.assertEqual(self.call(endpoint + '/batches/' + applied['batch_id'], 'POST', {}, owner)[0], 409)
        current = self.call(f'/construction/objects/{obj}', token=owner)[1]
        self.assertIn(14000, [item['unit_price'] for item in current['quantities']])

    def test_receipt_ocr_is_scoped_quota_bounded_and_only_prefills_a_draft(self):
        from PIL import Image

        owner, _ = self.account('ocr-owner')
        outsider, _ = self.account('ocr-outsider')
        obj = self.call('/construction/objects', 'POST', {'name': 'OCR object'}, owner)[1]['object']['id']
        other_obj = self.call('/construction/objects', 'POST', {'name': 'Other object'}, owner)[1]['object']['id']
        image = io.BytesIO()
        Image.new('RGB', (20, 20), (230, 230, 230)).save(image, format='PNG')
        file_id = self.call('/files', 'POST', {
            'construction_id': obj, 'name': 'receipt.png',
            'content': base64.b64encode(image.getvalue()).decode(),
        }, owner)[1]['file']['id']
        endpoint = f'/construction/objects/{obj}/receipt-ocr'
        with patch.dict(os.environ, OPENROUTER_API_KEY='test-only'), patch(
            'backend.receipt_ocr._request', side_effect=DomainError(429, 'provider busy')
        ):
            self.assertEqual(self.call(endpoint, 'POST', {'file_id': file_id}, owner)[0], 429)
        with patch.dict(os.environ, OPENROUTER_API_KEY='test-only', OPENROUTER_MODEL='paid-text-only'), patch(
            'backend.receipt_ocr._request',
            return_value={'merchant': 'Test store', 'amount_kopecks': 12500, 'date': '2026-09-30'},
        ) as model:
            self.assertEqual(self.call(endpoint, 'POST', {'file_id': file_id}, outsider)[0], 404)
            self.assertEqual(self.call(f'/construction/objects/{other_obj}/receipt-ocr', 'POST',
                                       {'file_id': file_id}, owner)[0], 400)
            for _ in range(3):
                code, result = self.call(endpoint, 'POST', {'file_id': file_id}, owner)
                self.assertEqual(code, 200, result)
                self.assertEqual(result['draft']['amount_kopecks'], 12500)
                self.assertTrue(result['needs_confirmation'])
            self.assertEqual(self.call(endpoint, 'POST', {'file_id': file_id}, owner)[0], 429)
            self.assertEqual(model.call_count, 3)
            self.assertTrue(all(call.args[0] == 'openrouter/free' for call in model.call_args_list))
        self.assertEqual(self.call(f'/construction/objects/{obj}', token=owner)[1]['purchases'], [])

    def test_construction_report_includes_photo_and_excludes_supplier_prices(self):
        owner, _ = self.account('report-owner')
        outsider, _ = self.account('report-outsider')
        _, created = self.call('/construction/objects', 'POST', {'name': 'Гостиная'}, owner)
        obj = created['object']['id']
        _, room = self.call(f'/construction/objects/{obj}/zones', 'POST',
                            {'name': 'Зал', 'length': '5', 'width': '4'}, owner)
        zone = room['zone']['id']
        _, work = self.call(f'/construction/objects/{obj}/quantities', 'POST', {
            'zone_id': zone, 'kind': 'work', 'title': 'Покраска', 'formula': 'area',
            'unit_price': 10000, 'cost_price': 77777,
        }, owner)
        work_id = work['quantity']['id']
        self.assertEqual(self.call(f'/construction/objects/{obj}/facts', 'POST',
                                   {'quantity_id': work_id, 'quantity': '4'}, owner)[0], 201)
        _, material = self.call(f'/construction/objects/{obj}/quantities', 'POST', {
            'zone_id': zone, 'kind': 'material', 'title': 'Краска', 'parent_work_id': work_id,
            'consumption_rate': '0.5', 'unit': 'л', 'unit_price': 2000,
        }, owner)
        _, supplier = self.call('/construction/suppliers', 'POST', {'name': 'Секретный поставщик'}, owner)
        self.assertEqual(self.call(f'/construction/objects/{obj}/purchases', 'POST', {
            'material_id': material['quantity']['id'], 'supplier_id': supplier['supplier']['id'],
            'purchased_on': '2026-09-30', 'quantity': '6', 'unit_price_kopecks': 1234567,
            'status': 'received',
        }, owner)[0], 201)
        from PIL import Image
        image = io.BytesIO()
        Image.new('RGB', (16, 12), (20, 70, 120)).save(image, format='PNG')
        _, photo = self.call('/files', 'POST', {
            'construction_id': obj, 'name': 'работа.png',
            'content': base64.b64encode(image.getvalue()).decode(),
        }, owner)
        self.assertEqual(self.call(f'/construction/objects/{obj}/logs', 'POST', {
            'work_date': '2026-09-30', 'work_description': 'Окрашена стена',
            'photo_file_ids': [photo['file']['id']],
        }, owner)[0], 201)
        status, payload = self.call(f'/construction/objects/{obj}/report.pdf', token=owner, raw=True)
        self.assertEqual(status, 200)
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(payload))
        report = '\n'.join(page.extract_text() for page in reader.pages)
        self.assertIn('Покраска', report)
        self.assertIn('Окрашена стена', report)
        self.assertNotIn('Секретный поставщик', report)
        self.assertNotIn('12 345,67', report)
        self.assertTrue(any(page.images for page in reader.pages))
        self.assertEqual(self.call(f'/construction/objects/{obj}/report.pdf', token=outsider)[0], 404)

    def test_additional_work_requires_separate_client_approval_and_preserves_versions(self):
        owner, _ = self.account('change-owner')
        outsider, _ = self.account('change-outsider')
        _, created = self.call('/construction/objects', 'POST', {'name': 'Квартира'}, owner)
        obj = created['object']['id']
        prefix = f'/construction/objects/{obj}/changes'
        payload = {'title': 'Дополнительные розетки', 'description': 'Три точки',
                   'items': [{'name': 'Розетка', 'quantity': '3', 'unit': 'шт.',
                              'unit_price': 150000, 'cost_price': 50000}], 'deadline_days': 2}
        status, created_change = self.call(prefix, 'POST', payload, owner)
        self.assertEqual(status, 201, created_change)
        change_id = created_change['change']['id']
        self.assertEqual(created_change['change']['amount_kopecks'], 450000)
        self.assertNotIn('cost_price', created_change['change']['items'][0])
        self.assertEqual(self.call(prefix + '/' + change_id + '/send', 'POST', {}, outsider)[0], 404)
        status, sent = self.call(prefix + '/' + change_id + '/send', 'POST', {}, owner)
        self.assertEqual(status, 200, sent)
        token = sent['change']['public_token']
        self.assertEqual(self.call(prefix + '/' + change_id, 'PATCH', {'title': 'Тихая правка'}, owner)[0], 409)
        public = self.call('/public/change?token=' + token)[1]['change']
        self.assertEqual(public['amount_kopecks'], 450000)
        self.assertNotIn('public_token', public)
        self.assertNotIn('cost_price', public['items'][0])
        self.assertEqual(self.call('/public/change/respond', 'POST',
                                   {'token': token, 'action': 'changes_requested', 'name': 'Клиент',
                                    'comment': 'Нужна другая розетка'})[0], 200)
        self.assertEqual(self.call('/public/change/respond', 'POST',
                                   {'token': token, 'action': 'approved', 'name': 'Клиент'})[0], 409)
        status, revised = self.call(prefix + '/' + change_id + '/revise', 'POST',
                                    {'items': [{'name': 'Розетка другая', 'quantity': '3',
                                                'unit': 'шт.', 'unit_price': 170000}]}, owner)
        self.assertEqual(status, 201, revised)
        self.assertEqual(revised['change']['version'], 2)
        self.assertEqual(revised['change']['amount_kopecks'], 510000)
        self.assertEqual(self.call(prefix + '/' + change_id + '/revise', 'POST', {}, owner)[0], 409)
        newer = revised['change']['id']
        token2 = self.call(prefix + '/' + newer + '/send', 'POST', {}, owner)[1]['change']['public_token']
        self.assertEqual(self.call('/public/change/respond', 'POST',
                                   {'token': token2, 'action': 'approved', 'name': 'Клиент'})[0], 200)
        detail = self.call(f'/construction/objects/{obj}', token=owner)[1]
        changes = detail['changes']
        self.assertEqual([row['status'] for row in changes], ['approved', 'changes_requested'])
        self.assertEqual(detail['scope']['approved_changes_kopecks'], 510000)
        status, extra = self.call(prefix, 'POST', payload, owner)
        self.assertEqual(status, 201, extra)
        extra_id = extra['change']['id']
        self.assertEqual(self.call(prefix + '/' + extra_id, 'DELETE', token=owner)[0], 200)
        status, extra = self.call(prefix, 'POST', payload, owner)
        self.assertEqual(status, 201, extra)
        extra_id = extra['change']['id']
        revoked_token = self.call(prefix + '/' + extra_id + '/send', 'POST', {}, owner)[1]['change']['public_token']
        self.assertEqual(self.call(prefix + '/' + extra_id + '/revoke', 'POST', {}, owner)[0], 200)
        self.assertEqual(self.call('/public/change?token=' + revoked_token)[0], 404)
        self.assertEqual(self.call(f'/construction/objects/{obj}', token=outsider)[0], 404)

    def test_construction_act_freezes_recorded_fact_instead_of_plan(self):
        owner, _ = self.account('act-owner')
        _, created = self.call('/construction/objects', 'POST', {'name': 'Квартира'}, owner)
        obj = created['object']['id']
        _, room = self.call(f'/construction/objects/{obj}/zones', 'POST',
                            {'name': 'Кухня', 'length': '5', 'width': '4'}, owner)
        _, work = self.call(f'/construction/objects/{obj}/quantities', 'POST', {
            'zone_id': room['zone']['id'], 'kind': 'work', 'title': 'Покраска',
            'formula': 'area', 'unit_price': 10000,
        }, owner)
        quantity = work['quantity']['id']
        self.assertEqual(self.call(f'/construction/objects/{obj}/act', 'POST', {}, owner)[0], 409)
        self.assertEqual(self.call(f'/construction/objects/{obj}/facts', 'POST',
                                   {'quantity_id': quantity, 'quantity': '3'}, owner)[0], 201)
        status, created_act = self.call(f'/construction/objects/{obj}/act', 'POST', {}, owner)
        self.assertEqual(status, 201, created_act)
        self.assertEqual(created_act['amount_kopecks'], 30000)
        document_id = created_act['document']['id']
        status, content = self.call('/documents/' + document_id + '/pdf', token=owner, raw=True)
        self.assertEqual(status, 200)
        from pypdf import PdfReader
        first = PdfReader(io.BytesIO(content)).pages[0].extract_text()
        self.assertIn('Покраска', first)
        self.assertIn('3 м²', first)
        self.assertNotIn('20 м²', first)
        self.assertEqual(self.call(f'/construction/objects/{obj}/facts', 'POST',
                                   {'quantity_id': quantity, 'quantity': '4'}, owner)[0], 201)
        status, unchanged = self.call('/documents/' + document_id + '/pdf', token=owner, raw=True)
        self.assertEqual(status, 200)
        self.assertEqual(first, PdfReader(io.BytesIO(unchanged)).pages[0].extract_text())

    def test_procurement_tracks_received_material_and_scopes_supplier_and_receipt(self):
        owner, _ = self.account('purchase-owner')
        outsider, _ = self.account('purchase-outsider')
        _, created = self.call('/construction/objects', 'POST', {'name': 'Ремонт'}, owner)
        obj = created['object']['id']
        _, room = self.call(f'/construction/objects/{obj}/zones', 'POST', {
            'name': 'Кухня', 'length': '5', 'width': '4',
        }, owner)
        zone = room['zone']['id']
        _, work = self.call(f'/construction/objects/{obj}/quantities', 'POST', {
            'zone_id': zone, 'kind': 'work', 'title': 'Покраска', 'formula': 'area', 'unit_price': 1000,
        }, owner)
        _, material = self.call(f'/construction/objects/{obj}/quantities', 'POST', {
            'zone_id': zone, 'kind': 'material', 'title': 'Краска', 'unit': 'л',
            'parent_work_id': work['quantity']['id'], 'consumption_rate': '0.5', 'unit_price': 2000,
        }, owner)
        material_id = material['quantity']['id']
        _, supplier = self.call('/construction/suppliers', 'POST', {'name': 'Склад', 'phone': '+70000000000'}, owner)
        supplier_id = supplier['supplier']['id']
        from PIL import Image
        image = io.BytesIO()
        Image.new('RGB', (4, 4), (20, 30, 40)).save(image, format='PNG')
        _, uploaded = self.call('/files', 'POST', {
            'construction_id': obj, 'name': 'чек.png',
            'content': base64.b64encode(image.getvalue()).decode(),
        }, owner)
        receipt_id = uploaded['file']['id']
        payload = dict(material_id=material_id, supplier_id=supplier_id, receipt_file_id=receipt_id,
                       purchased_on='2026-09-30', quantity='6', unit_price_kopecks=10000, status='received')
        status, purchase = self.call(f'/construction/objects/{obj}/purchases', 'POST', payload, owner)
        self.assertEqual(status, 201, purchase)
        purchase_id = purchase['purchase']['id']
        detail = self.call(f'/construction/objects/{obj}', token=owner)[1]
        self.assertEqual(detail['procurement'][0]['required_quantity'], '10.0000')
        self.assertEqual(detail['procurement'][0]['received_quantity'], '6')
        self.assertEqual(detail['procurement'][0]['to_buy_quantity'], '4.0000')
        self.assertEqual(self.call(f'/construction/objects/{obj}/purchases/{purchase_id}', 'PATCH',
                                   {'status': 'ordered'}, outsider)[0], 404)
        self.assertEqual(self.call('/construction/suppliers/' + supplier_id, 'PATCH',
                                   {'name': 'Чужой'}, outsider)[0], 404)
        self.assertEqual(self.call(f'/construction/objects/{obj}/purchases', 'POST',
                                   dict(payload, quantity='-1'), owner)[0], 400)
        self.assertEqual(self.call(f'/construction/objects/{obj}/purchases/{purchase_id}', 'PATCH',
                                   {'quantity': '12'}, owner)[0], 200)
        detail = self.call(f'/construction/objects/{obj}', token=owner)[1]
        self.assertEqual(detail['procurement'][0]['to_buy_quantity'], '0')
        self.assertEqual(self.call(f'/construction/objects/{obj}/purchases/{purchase_id}', 'DELETE',
                                   token=owner)[0], 200)
        self.assertEqual(self.call(f'/construction/objects/{obj}', token=owner)[1]['procurement'][0]['received_quantity'], '0')

    def test_daily_log_tracks_work_photo_and_fact_atomically(self):
        owner, _ = self.account('log-owner')
        outsider, _ = self.account('log-outsider')
        _, created = self.call('/construction/objects', 'POST', {'name': 'Ремонт'}, owner)
        obj = created['object']['id']
        _, room = self.call(f'/construction/objects/{obj}/zones', 'POST', {
            'name': 'Зал', 'length': '5', 'width': '4',
        }, owner)
        zone_id = room['zone']['id']
        _, work = self.call(f'/construction/objects/{obj}/quantities', 'POST', {
            'zone_id': zone_id, 'title': 'Штукатурка', 'formula': 'area',
            'unit': 'м²', 'unit_price': 5000,
        }, owner)
        quantity_id = work['quantity']['id']
        from PIL import Image

        image = io.BytesIO()
        Image.new('RGB', (4, 4), (30, 40, 50)).save(image, format='PNG')
        _, uploaded = self.call('/files', 'POST', {
            'construction_id': obj, 'name': 'работа.png',
            'content': base64.b64encode(image.getvalue()).decode(),
        }, owner)
        photo_id = uploaded['file']['id']
        payload = {
            'work_date': '2026-09-30', 'zone_id': zone_id, 'quantity_id': quantity_id,
            'workers': 'Иван и Анна', 'worker_count': 2,
            'work_description': 'Штукатурка стены', 'completed_quantity': '8',
            'comment': 'Первый слой', 'photo_file_ids': [photo_id],
        }
        status, created_log = self.call(f'/construction/objects/{obj}/logs', 'POST', payload, owner)
        self.assertEqual(status, 201, created_log)
        log_id = created_log['daily_log']['id']
        fact_id = created_log['daily_log']['fact_id']
        self.assertEqual(created_log['daily_log']['photo_file_ids'], [photo_id])
        detail = self.call(f'/construction/objects/{obj}', token=owner)[1]
        self.assertEqual(detail['quantities'][0]['actual_quantity'], '8')
        self.assertEqual(detail['daily_logs'][0]['photo_file_ids'], [photo_id])
        self.assertEqual(self.call(f'/construction/objects/{obj}/facts/{fact_id}', 'DELETE', token=owner)[0], 409)
        self.assertEqual(self.call(f'/construction/objects/{obj}/logs', 'POST',
                                   dict(payload, photo_file_ids=[photo_id, photo_id]), owner)[0], 400)
        self.assertEqual(self.call(f'/construction/objects/{obj}/logs/{log_id}', 'PATCH', {
            'completed_quantity': '10', 'comment': 'Уточнённый объём',
        }, owner)[0], 200)
        detail = self.call(f'/construction/objects/{obj}', token=owner)[1]
        self.assertEqual(detail['quantities'][0]['actual_quantity'], '10')
        self.assertEqual(self.call(f'/construction/objects/{obj}/logs/{log_id}', 'PATCH', {
            'completed_quantity': '12',
        }, outsider)[0], 404)
        self.assertEqual(self.call(f'/construction/objects/{obj}/logs/{log_id}', 'DELETE', token=owner)[0], 200)
        detail = self.call(f'/construction/objects/{obj}', token=owner)[1]
        self.assertEqual(detail['quantities'][0]['actual_quantity'], '0')
        self.assertEqual(detail['daily_logs'], [])

    def test_chat_upload_is_private_searchable_and_limited_to_readable_files(self):
        owner, _ = self.account('chat-file-owner')
        other, _ = self.account('chat-file-other')
        payload = {'assistant_upload': True, 'name': 'inspection.md',
                   'content': base64.b64encode(b'# Inspection\nPaint sample').decode()}
        status, uploaded = self.call('/files', 'POST', payload, owner)
        self.assertEqual(status, 201, uploaded)
        file_id = uploaded['file']['id']
        self.assertEqual(self.call('/files/' + file_id, token=owner, raw=True)[1], b'# Inspection\nPaint sample')
        self.assertEqual(self.call('/files/' + file_id, token=other)[0], 404)
        self.assertEqual(self.call('/search?q=Paint%20sample', token=owner)[1]['items'][0]['file_id'], file_id)
        self.assertEqual(self.call('/files', 'POST', dict(payload, public=1), owner)[0], 400)
        self.assertEqual(self.call('/files', 'POST', dict(payload, name='inspection.png'), owner)[0], 400)

    def test_measurement_change_recalculates_work_without_rewriting_quote(self):
        token, _ = self.account('measure-owner')
        other, _ = self.account('measure-other')
        _, created = self.call('/construction/objects', 'POST', {'name': 'Замеры'}, token)
        obj = created['object']['id']
        _, zone_result = self.call(f'/construction/objects/{obj}/zones', 'POST', {
            'name': 'Кухня', 'length': '4', 'width': '3',
        }, token)
        _, measure_result = self.call(f'/construction/objects/{obj}/measurements', 'POST', {
            'zone_id': zone_result['zone']['id'], 'symbol': 'beam_length',
            'value': '200', 'unit': 'см',
        }, token)
        measure_id = measure_result['measurement']['id']
        _, quantity_result = self.call(f'/construction/objects/{obj}/quantities', 'POST', {
            'zone_id': zone_result['zone']['id'], 'title': 'Балка',
            'formula': 'beam_length * 2', 'unit': 'м', 'unit_price': 5000,
        }, token)
        self.assertEqual(quantity_result['quantity']['quantity'], '4.0000')
        _, old = self.call(f'/construction/objects/{obj}/quote', 'POST', {}, token)
        status, changed = self.call(f'/construction/objects/{obj}/measurements/{measure_id}', 'PATCH', {
            'value': '3', 'unit': 'м',
        }, token)
        self.assertEqual(status, 200, changed)
        self.assertEqual(changed['quantities'][0]['quantity'], '6.0000')
        self.assertIsNone(changed['object']['quote_id'])
        self.assertEqual(self.call('/quotes/' + old['quote']['id'], token=token)[1]['quote']['amount_kopecks'], 20000)
        self.assertEqual(self.call(f'/construction/objects/{obj}/measurements/{measure_id}', 'PATCH', {
            'value': '5',
        }, other)[0], 404)

    def test_markdown_file_edit_is_versioned_scoped_and_reindexed(self):
        owner, _ = self.account('markdown-owner')
        other, _ = self.account('markdown-other')
        _, client = self.call('/clients', 'POST', {'name': 'Контакт'}, owner)
        status, result = self.call('/files', 'POST', {
            'client_id': client['item']['id'], 'name': 'замеры.md',
            'content': base64.b64encode(b'# Plan\nOld term').decode(),
        }, owner)
        self.assertEqual(status, 201, result)
        file = result['file']
        path = '/files/' + file['id']
        status, _ = self.call(path, 'PATCH', {
            'sha256': file['sha256'],
            'content': base64.b64encode(b'# Plan\nNew term').decode(),
        }, other)
        self.assertEqual(status, 404)
        status, changed = self.call(path, 'PATCH', {
            'sha256': file['sha256'],
            'content': base64.b64encode(b'# Plan\nNew term').decode(),
        }, owner)
        self.assertEqual(status, 200, changed)
        self.assertNotEqual(changed['file']['sha256'], file['sha256'])
        self.assertEqual(self.call(path, token=owner, raw=True)[1], b'# Plan\nNew term')
        self.assertEqual(self.call(path, 'PATCH', {
            'sha256': file['sha256'],
            'content': base64.b64encode(b'# Plan\nAnother term').decode(),
        }, owner)[0], 409)
        self.assertEqual(len(self.call('/search?q=New%20term', token=owner)[1]['items']), 1)
        history = self.call(path + '/versions', token=owner)[1]['items']
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]['revision'], 1)
        self.assertEqual(self.call(path + '/versions', token=other)[0], 404)
        version_path = path + '/versions/' + history[0]['id']
        self.assertEqual(
            base64.b64decode(self.call(version_path, token=owner)[1]['content']),
            b'# Plan\nOld term',
        )
        self.assertEqual(self.call(version_path, token=other)[0], 404)
        self.assertEqual(self.call(version_path + '/restore', 'POST', {
            'sha256': file['sha256'],
        }, owner)[0], 409)
        restored = self.call(version_path + '/restore', 'POST', {
            'sha256': changed['file']['sha256'],
        }, owner)
        self.assertEqual(restored[0], 200, restored)
        self.assertEqual(self.call(path, token=owner, raw=True)[1], b'# Plan\nOld term')
        self.assertEqual(len(self.call(path + '/versions', token=owner)[1]['items']), 2)
        self.assertEqual(len(self.call('/search?q=New%20term', token=owner)[1]['items']), 0)

    def test_construction_measure_to_quote_and_workspace_isolation(self):
        token, _ = self.account('builder-' + os.urandom(4).hex())
        other, _ = self.account('other-builder-' + os.urandom(4).hex())
        status, result = self.call('/construction/objects', 'POST', {'name': 'Квартира'}, token)
        self.assertEqual(status, 201, result)
        obj = result['object']['id']
        status, result = self.call(f'/construction/objects/{obj}/zones', 'POST', {
            'name': 'Кухня', 'length': '5', 'width': '4', 'height': '2.8',
        }, token)
        self.assertEqual(status, 201, result)
        zone = result['zone']['id']
        self.assertEqual(result['calculated']['area'], '20')
        status, result = self.call(f'/construction/objects/{obj}/quantities', 'POST', {
            'zone_id': zone, 'title': 'Покраска', 'formula': 'area', 'unit_price': 15000,
        }, token)
        self.assertEqual(status, 201, result)
        work = result['quantity']
        self.assertEqual(work['planned_total_kopecks'], 300000)
        status, result = self.call(f'/construction/objects/{obj}/quantities', 'POST', {
            'parent_work_id': work['id'], 'kind': 'material', 'title': 'Краска',
            'unit': 'л', 'consumption_rate': '0.2', 'waste_percent': 10,
            'unit_price': 80000,
        }, token)
        self.assertEqual(status, 201, result)
        self.assertEqual(result['quantity']['quantity'], '4.4000')
        status, result = self.call(f'/construction/objects/{obj}/facts', 'POST', {
            'quantity_id': work['id'], 'quantity': '8',
        }, token)
        self.assertEqual(status, 201, result)
        status, result = self.call(f'/construction/objects/{obj}', token=token)
        self.assertEqual(status, 200, result)
        self.assertEqual(result['totals']['actual_kopecks'], 120000)
        status, result = self.call(f'/construction/objects/{obj}/quote', 'POST', {}, token)
        self.assertEqual(status, 201, result)
        self.assertEqual(result['quote']['amount_kopecks'], 652000)
        status, recalculated = self.call(f'/construction/objects/{obj}/zones/{zone}', 'PATCH', {
            'width': '480', 'dimension_unit': 'см',
        }, token)
        self.assertEqual(status, 200, recalculated)
        self.assertIsNone(recalculated['object']['quote_id'])
        amounts = {row['id']: row['quantity'] for row in recalculated['quantities']}
        self.assertEqual(amounts[work['id']], '24.0000')
        self.assertIn('5.2800', amounts.values())
        status, updated = self.call(f'/construction/objects/{obj}/quantities/{work["id"]}', 'PATCH', {
            'unit_price': 16000,
        }, token)
        self.assertEqual(status, 200, updated)
        self.assertEqual(next(row for row in updated['quantities'] if row['id'] == work['id'])['planned_total_kopecks'], 384000)
        status, priced_plan = self.call(f'/construction/objects/{obj}/quantities/{work["id"]}', 'PATCH', {
            'price_coefficient': '1.2', 'markup_percent': '15', 'discount_percent': '10',
            'coefficient_reason': 'Сложная поверхность',
        }, token)
        self.assertEqual(status, 200, priced_plan)
        self.assertEqual(next(row for row in priced_plan['quantities'] if row['id'] == work['id'])['planned_total_kopecks'], 476928)
        status, fresh_quote = self.call(f'/construction/objects/{obj}/quote', 'POST', {}, token)
        self.assertEqual(status, 201, fresh_quote)
        self.assertEqual(fresh_quote['quote']['amount_kopecks'], priced_plan['totals']['planned_kopecks'])
        published = self.publish(token, fresh_quote['quote'])
        public_token = published['public_url'].split('quote=')[1]
        self.assertEqual(self.call('/public/accept', 'POST', {
            'token': public_token, 'version': 1, 'name': 'Клиент',
        })[0], 200)
        status, project = self.call('/quotes/' + published['id'] + '/project', 'POST', {}, token)
        self.assertEqual(status, 201, project)
        self.assertEqual(self.call(f'/construction/objects/{obj}', token=token)[1]['object']['project_id'], project['project']['id'])
        status, sheet = self.call(f'/construction/objects/{obj}/xlsx', token=token, raw=True)
        self.assertEqual(status, 200)
        self.assertTrue(sheet.startswith(b'PK'))
        status, result = self.call('/construction/objects', 'POST', {'name': 'Импорт'}, token)
        self.assertEqual(status, 201, result)
        imported_object = result['object']['id']
        status, result = self.call(f'/construction/objects/{imported_object}/zones', 'POST', {
            'name': 'Кухня', 'length': '5', 'width': '4.8',
        }, token)
        self.assertEqual(status, 201, result)
        status, result = self.call(f'/construction/objects/{imported_object}/import', 'POST', {
            'content': base64.b64encode(sheet).decode(),
        }, token)
        self.assertEqual(status, 201, result)
        self.assertEqual(result['imported'], 2)
        status, imported = self.call(f'/construction/objects/{imported_object}', token=token)
        self.assertEqual(status, 200, imported)
        self.assertEqual(len(imported['quantities']), 2)
        from openpyxl import load_workbook

        malicious = load_workbook(io.BytesIO(sheet))
        malicious.active.cell(row=2, column=2).value = '=2+2'
        output = io.BytesIO()
        malicious.save(output)
        status, _ = self.call(f'/construction/objects/{imported_object}/import', 'POST', {
            'content': base64.b64encode(output.getvalue()).decode(),
        }, token)
        self.assertEqual(status, 400)
        self.assertEqual(len(self.call(f'/construction/objects/{imported_object}', token=token)[1]['quantities']), 2)
        status, result = self.call('/files', 'POST', {
            'construction_id': obj, 'name': 'замер.txt',
            'content': base64.b64encode('Кухня 5 × 4 м'.encode()).decode(),
        }, token)
        self.assertEqual(status, 201, result)
        status, result = self.call('/files?construction_id=' + obj, token=token)
        self.assertEqual(status, 200, result)
        self.assertEqual(len(result['items']), 1)
        status, defect = self.call(f'/construction/objects/{obj}/defects', 'POST', {
            'zone_id': zone, 'description': 'Трещина в углу',
            'measurement_note': '1,2 м', 'suggested_work': 'Заделка', 'severity': 'high',
        }, token)
        self.assertEqual(status, 201, defect)
        defect_id = defect['defect']['id']
        self.assertEqual(self.call(f'/construction/objects/{obj}/defects/{defect_id}', 'PATCH', {
            'status': 'resolved',
        }, token)[1]['defect']['status'], 'resolved')
        self.assertEqual(self.call(f'/construction/objects/{obj}', token=token)[1]['defects'][0]['id'], defect_id)
        self.assertEqual(self.call(f'/construction/objects/{obj}/defects', 'POST', {
            'description': 'Неверное фото', 'photo_file_id': result['items'][0]['id'],
        }, token)[0], 400)
        from PIL import Image

        photo = io.BytesIO()
        Image.new('RGB', (4, 4), (30, 80, 140)).save(photo, format='PNG')
        status, uploaded = self.call('/files', 'POST', {
            'construction_id': obj, 'name': 'дефект.png',
            'content': base64.b64encode(photo.getvalue()).decode(),
        }, token)
        self.assertEqual(status, 201, uploaded)
        photo_id = uploaded['file']['id']
        status, linked = self.call(f'/construction/objects/{obj}/defects/{defect_id}', 'PATCH', {
            'photo_file_id': photo_id,
        }, token)
        self.assertEqual(status, 200, linked)
        self.assertEqual(linked['defect']['photo_file_id'], photo_id)
        self.assertEqual(self.call(f'/construction/objects/{imported_object}/defects', 'POST', {
            'description': 'Чужое фото', 'photo_file_id': photo_id,
        }, token)[0], 400)
        self.assertEqual(self.call(f'/construction/objects/{imported_object}/defects/{defect_id}', 'PATCH', {
            'status': 'open',
        }, token)[0], 404)
        status, _ = self.call(f'/construction/objects/{obj}', token=other)
        self.assertEqual(status, 404)
        status, _ = self.call('/files?construction_id=' + obj, token=other)
        self.assertEqual(status, 404)

    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.environ = dict(os.environ)
        os.environ["UPLOAD_DIR"] = str(Path(cls.temp.name) / "uploads")
        os.environ.update(
            DB_PATH=str(Path(cls.temp.name) / "business.sqlite3"),
            PORT="0",
            PUBLIC_ORIGIN="http://localhost:0",
        )
        os.environ.pop("SMTP_HOST", None)
        spec = importlib.util.spec_from_file_location(
            "business_app", Path(__file__).resolve().parents[1] / "backend/app.py"
        )
        cls.mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.mod)
        cls.mod.migrate()
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), cls.mod.Handler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.base = "http://127.0.0.1:" + str(cls.server.server_address[1])

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.temp.cleanup()
        os.environ.clear()
        os.environ.update(cls.environ)

    def setUp(self):
        self.mod.RATE.clear()

    def call(
        self,
        path,
        method="GET",
        data=None,
        token=None,
        workspace=None,
        key=None,
        raw=False,
    ):
        headers = {"Content-Type": "application/json"}
        if token:
            headers["Authorization"] = "Bearer " + token
        if workspace:
            headers["X-Workspace-Id"] = workspace
        if key:
            headers["Idempotency-Key"] = key
        req = urllib.request.Request(
            self.base + "/api" + path,
            json.dumps(data).encode() if data is not None else None,
            headers,
            method=method,
        )
        try:
            with urllib.request.urlopen(req) as response:
                return response.status, response.read() if raw else json.load(response)
        except urllib.error.HTTPError as response:
            with response:
                return response.code, json.load(response)

    def account(self, name):
        code, result = self.call(
            "/auth/register",
            "POST",
            {
                "email": name + "@test.invalid",
                "name": name,
                "password": "secure twelve password",
            },
        )
        self.assertEqual(code, 200, result)
        return result["token"], result["user"]["id"]

    def test_quote_creation_key_replays_atomically_and_counts_once(self):
        token, user_id = self.account("quote-replay-owner")
        body = {"title": "Повтор при плохой сети", "client": "Клиент", "amount": 10000}
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(
                lambda _: self.call("/quotes", "POST", body, token, key="quote-retry"), range(2)
            ))
        self.assertEqual(sorted(code for code, _ in results), [200, 201], results)
        self.assertEqual(results[0][1]["quote"]["id"], results[1][1]["quote"]["id"])
        self.assertEqual(len(self.call("/quotes", token=token)[1]["quotes"]), 1)
        with self.mod.db() as con:
            count = con.execute(
                "SELECT count(*) FROM events WHERE user_id=? AND name='quote_created'", (user_id,)
            ).fetchone()[0]
        self.assertEqual(count, 1)
        status, replay = self.call("/quotes", "POST", {**body, "_request_key": "client-only"}, token, key="quote-retry")
        self.assertEqual(status, 200, replay)
        self.assertTrue(replay["replayed"])
        self.assertEqual(self.call("/quotes", "POST", {**body, "amount": 20000}, token, key="quote-retry")[0], 409)

    def test_quote_creation_key_is_isolated_and_does_not_recreate_deleted_quote(self):
        owner, _ = self.account("quote-key-owner")
        outsider, _ = self.account("quote-key-outsider")
        body = {"title": "Смета", "client": "Клиент", "amount": 10000}
        first = self.call("/quotes", "POST", body, owner, key="shared-client-key")[1]["quote"]
        status, other = self.call("/quotes", "POST", body, outsider, key="shared-client-key")
        self.assertEqual(status, 201, other)
        self.assertNotEqual(first["id"], other["quote"]["id"])
        self.assertEqual(self.call("/quotes/" + first["id"], "DELETE", token=owner)[0], 200)
        self.assertEqual(self.call("/quotes", "POST", body, owner, key="shared-client-key")[0], 409)
        self.assertEqual(self.call("/quotes", token=owner)[1]["quotes"], [])

    def test_failed_quote_request_does_not_claim_key_and_replay_ignores_quota(self):
        token, _ = self.account("quote-key-quota")
        body = {"title": "Смета", "client": "Клиент", "amount": 10000}
        self.assertEqual(self.call("/quotes", "POST", {**body, "title": ""}, token, key="retry-validation")[0], 400)
        status, first = self.call("/quotes", "POST", body, token, key="retry-validation")
        self.assertEqual(status, 201, first)
        for _ in range(9):
            self.assertEqual(self.call("/quotes", "POST", body, token)[0], 201)
        self.assertEqual(self.call("/quotes", "POST", body, token)[0], 402)
        status, replay = self.call("/quotes", "POST", body, token, key="retry-validation")
        self.assertEqual(status, 200, replay)
        self.assertEqual(replay["quote"]["id"], first["quote"]["id"])

    def test_quote_catalog_usage_is_atomic_scoped_and_not_replayed(self):
        owner, _ = self.account("quote-usage-owner")
        outsider, _ = self.account("quote-usage-outsider")
        owned = [self.call("/catalog", "POST", {"name": name, "price": 10000}, owner)[1]["item"]["id"]
                 for name in ("Работа", "Материал")]
        foreign = self.call("/catalog", "POST", {"name": "Чужая работа", "price": 10000}, outsider)[1]["item"]["id"]
        body = {"title": "Смета из расценок", "client": "Клиент", "amount": 10000,
                "catalog_ids": [*owned, owned[0], foreign, "deleted-catalog-item"]}
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(
                lambda _: self.call("/quotes", "POST", body, owner, key="catalog-retry"), range(2)))
        self.assertEqual(sorted(code for code, _ in results), [200, 201], results)
        recent = self.call("/catalog?recent=1", token=owner)[1]["items"]
        self.assertEqual({item["id"] for item in recent}, set(owned))
        self.assertTrue(all(item["usage_count"] == 1 and item["last_used_at"] > 0 for item in recent))
        self.assertEqual(self.call("/catalog?recent=1", token=outsider)[1]["items"], [])

    def test_quote_catalog_usage_validation_leaves_no_quote_or_usage(self):
        owner, _ = self.account("quote-usage-validation")
        item = self.call("/catalog", "POST", {"name": "Работа", "price": 10000}, owner)[1]["item"]
        body = {"title": "Смета", "client": "Клиент", "amount": 10000}
        for catalog_ids in (None, "invalid", [None], [item["id"]] * 201):
            self.assertEqual(self.call("/quotes", "POST", {**body, "catalog_ids": catalog_ids},
                                       owner, key="validated-usage")[0], 400)
        self.assertEqual(self.call("/quotes", token=owner)[1]["quotes"], [])
        self.assertEqual(self.call("/catalog?recent=1", token=owner)[1]["items"], [])
        self.assertEqual(self.call("/quotes", "POST", {**body, "catalog_ids": [item["id"]]},
                                   owner, key="validated-usage")[0], 201)
        self.assertEqual(self.call("/catalog?recent=1", token=owner)[1]["items"][0]["usage_count"], 1)

    def quote(self, token, **extra):
        data = {
            "title": "Project",
            "client": "Client",
            "items": [
                {
                    "name": "Service",
                    "quantity": "2.5",
                    "unit": "hour",
                    "unit_price": 12000,
                    "cost_price": 4000,
                }
            ],
        }
        data.update(extra)
        code, result = self.call("/quotes", "POST", data, token)
        self.assertEqual(code, 201, result)
        return result["quote"]

    def publish(self, token, quote):
        code, result = self.call(
            "/quotes/" + quote["id"] + "/publish",
            "POST",
            {"revision": quote["revision"]},
            token,
        )
        self.assertEqual(code, 200, result)
        return result["quote"]

    def test_catalog_price_history_filters_favorites_and_workspace_scope(self):
        owner, _ = self.account("catalog-history-owner")
        outsider, _ = self.account("catalog-history-outsider")
        code, created = self.call("/catalog", "POST", {
            "name": "Покраска стен", "category": "Отделка", "article": "PT-1",
            "unit": "м²", "price": 12000, "cost_price": 7000,
        }, owner)
        self.assertEqual(code, 201, created)
        item = created["item"]
        item_id = item["id"]
        self.assertEqual(self.call("/catalog/" + item_id, token=outsider)[0], 404)
        self.assertEqual(self.call("/catalog?favorite=1", token=owner)[1]["items"], [])
        self.assertEqual(self.call("/catalog?q=PT-1", token=owner)[1]["items"][0]["id"], item_id)
        category = urllib.parse.quote("Отделка")
        self.assertEqual(self.call("/catalog?category=" + category, token=owner)[1]["categories"], ["Отделка"])

        code, saved = self.call("/catalog/" + item_id, "PATCH", {
            "revision": item["revision"], "favorite": 1,
        }, owner)
        self.assertEqual(code, 200, saved)
        self.assertEqual(self.call("/catalog?favorite=1", token=owner)[1]["items"][0]["id"], item_id)
        detail = self.call("/catalog/" + item_id, token=owner)[1]["item"]
        self.assertEqual([row["price"] for row in detail["price_history"]], [12000])

        self.assertEqual(self.call("/catalog/" + item_id, "PATCH", {
            "revision": item["revision"], "price": 15000,
        }, owner)[0], 409)
        code, saved = self.call("/catalog/" + item_id, "PATCH", {
            "revision": detail["revision"], "price": 15000,
        }, owner)
        self.assertEqual(code, 200, saved)
        detail = self.call("/catalog/" + item_id, token=owner)[1]["item"]
        self.assertEqual([row["price"] for row in detail["price_history"]], [15000, 12000])
        self.assertEqual(self.call("/catalog?category=" + urllib.parse.quote("Другое"), token=owner)[1]["items"], [])
        self.assertEqual(self.call("/catalog?category=" + category, token=outsider)[1]["items"], [])
        self.assertEqual(self.call("/catalog/" + item_id + "/use", "POST", {}, outsider)[0], 404)
        self.assertEqual(self.call("/catalog?recent=1", token=owner)[1]["items"], [])
        self.assertEqual(self.call("/catalog/" + item_id + "/use", "POST", {}, owner)[0], 200)
        recent = self.call("/catalog?recent=1", token=owner)[1]["items"]
        self.assertEqual(recent[0]["id"], item_id)
        self.assertEqual(recent[0]["usage_count"], 1)
        self.assertEqual(len(self.call("/catalog/" + item_id, token=owner)[1]["item"]["price_history"]), 2)
        duplicate_csv = "name,price,unit\nПокраска стен,16000,м²"
        self.assertEqual(self.call("/transfer/catalog", "POST", {"csv": duplicate_csv}, owner)[0], 409)
        self.assertEqual(len(self.call("/catalog", token=owner)[1]["items"]), 1)
        self.assertEqual(self.call("/transfer/catalog", "POST", {
            "csv": duplicate_csv, "allow_duplicates": True,
        }, owner)[0], 201)
        self.assertEqual(self.call("/catalog/" + item_id, "DELETE", token=owner)[0], 200)
        with self.mod.db() as con:
            self.assertEqual(con.execute(
                "SELECT count(*) FROM catalog_price_history WHERE item_id=?", (item_id,)
            ).fetchone()[0], 0)

    def test_dashboard_returns_five_lightweight_quotes_in_own_workspace(self):
        token, _ = self.account("dashboard-owner")
        other_token, _ = self.account("dashboard-other")
        for number in range(6):
            self.quote(token, title=f"Estimate {number}")
        self.quote(other_token, title="Private estimate")

        code, dashboard = self.call("/dashboard", token=token)
        self.assertEqual(code, 200)
        self.assertEqual(dashboard["workspace"]["role"], "owner")
        self.assertEqual(dashboard["overview"]["quotes"]["total"], 6)
        self.assertEqual(len(dashboard["quotes"]), 5)
        self.assertTrue(all(quote["title"].startswith("Estimate ") for quote in dashboard["quotes"]))
        self.assertTrue(all("items" not in quote for quote in dashboard["quotes"]))
        self.assertTrue(all("public_token" not in quote for quote in dashboard["quotes"]))
        self.assertNotIn("Private estimate", [q["title"] for q in dashboard["quotes"]])
        self.assertIn("ai_drafting", dashboard["capabilities"])
        code, full_list = self.call("/quotes", token=token)
        self.assertEqual(code, 200)
        self.assertEqual(len(full_list["quotes"]), 6)
        self.assertTrue(all(len(quote["items"]) == 1 for quote in full_list["quotes"]))
        self.assertEqual(self.call("/dashboard")[0], 401)

        with self.mod.db() as con:
            user_id = dashboard["workspace"]["owner_id"]
            handler = SimpleNamespace(
                headers={},
                auth=lambda connection: connection.execute(
                    "SELECT * FROM users WHERE id=?", (user_id,)
                ).fetchone(),
            )
            statements = []
            con.set_trace_callback(statements.append)
            Service(handler, con, self.base).handle("GET", "/api/quotes", {})
            item_reads = [
                sql for sql in statements if sql.lstrip().upper().startswith("SELECT")
                and "quote_items" in sql
            ]
            self.assertEqual(len(item_reads), 1)
            statements.clear()
            Service(handler, con, self.base).handle("GET", "/api/dashboard", {})
            self.assertFalse(any("quote_items" in sql for sql in statements))

    def test_today_lists_real_followups_without_other_workspace_data(self):
        token, _ = self.account("today-owner")
        other_token, _ = self.account("today-other")
        quote = self.publish(token, self.quote(token, title="Approved work"))
        public_token = quote["public_url"].split("quote=")[1]
        self.assertEqual(self.call("/public/accept", "POST", {
            "token": public_token, "version": 1, "name": "Client"
        })[0], 200)
        self.quote(other_token, title="Private work")
        yesterday = (dt.date.today() - dt.timedelta(days=1)).isoformat()
        task_status, task = self.call("/tasks", "POST", {
            "name": "Call the client", "due_date": yesterday
        }, token)
        self.assertEqual(task_status, 201, task)

        code, dashboard = self.call("/dashboard", token=token)
        self.assertEqual(code, 200, dashboard)
        actions = dashboard["actions"]
        self.assertEqual(actions["summary"]["approved_this_month"], 1)
        self.assertEqual(actions["summary"]["active_projects"], 0)
        self.assertIn(quote["id"], [a["entity_id"] for a in actions["items"]])
        self.assertIn(task["item"]["id"], [a["entity_id"] for a in actions["items"]])
        self.assertNotIn("Private work", json.dumps(actions))

        code, project = self.call("/quotes/" + quote["id"] + "/project", "POST", {}, token)
        self.assertEqual(code, 201, project)
        _, dashboard = self.call("/dashboard", token=token)
        self.assertEqual(dashboard["actions"]["summary"]["active_projects"], 1)
        self.assertEqual(dashboard["actions"]["summary"]["waiting_payments"], 1)
        self.assertNotIn(quote["id"], [a["entity_id"] for a in dashboard["actions"]["items"]])

    def test_public_intake_creates_private_client_lead_request_and_attachment(self):
        owner, _ = self.account("intake-owner")
        outsider, _ = self.account("intake-outsider")
        code, result = self.call("/intake", "POST", {}, owner)
        self.assertEqual(code, 200, result)
        form = result["form"]
        token = form["public_url"].split("intake=")[1]
        self.assertEqual(self.call("/public/intake?token=" + token)[0], 200)
        payload = {
            "token": token, "name": "Ирина", "email": "irina@example.ru",
            "details": "Нужен фирменный сайт", "budget_kopecks": 7500000,
            "due_date": "2026-12-15", "comment": "Напишите после обеда",
            "file": {"name": "brief.txt", "content": base64.b64encode(b"Project brief").decode()},
        }
        self.assertEqual(self.call("/public/intake", "POST", {**payload, "email": ""})[0], 400)
        self.assertEqual(self.call("/public/intake", "POST", {**payload, "file": {
            "name": "active.html", "content": base64.b64encode(b"<script></script>").decode()
        }})[0], 400)
        self.assertEqual(self.call("/leads", token=owner)[1]["items"], [])
        code, submitted = self.call("/public/intake", "POST", payload)
        self.assertEqual(code, 201, submitted)
        self.assertEqual(submitted, {"ok": True})
        _, leads = self.call("/leads", token=owner)
        self.assertEqual(len(leads["items"]), 1)
        lead_id = leads["items"][0]["id"]
        _, today = self.call("/dashboard", token=owner)
        self.assertIn(lead_id, [action["entity_id"] for action in today["actions"]["items"]])
        _, detail = self.call("/leads/" + lead_id, token=owner)
        self.assertEqual(detail["item"]["request"]["details"], payload["details"])
        self.assertEqual(detail["item"]["request"]["budget_kopecks"], 7500000)
        client_id = detail["item"]["request"]["client_id"]
        _, client = self.call("/clients/" + client_id, token=owner)
        self.assertEqual(len(client["item"]["requests"]), 1)
        self.assertTrue(any(event["action"] == "Новая заявка" for event in client["item"]["timeline"]))
        _, files = self.call("/files?client_id=" + client_id, token=owner)
        self.assertEqual([file["name"] for file in files["items"]], ["brief.txt"])
        self.assertEqual(self.call("/leads/" + lead_id, token=outsider)[0], 404)
        self.assertEqual(self.call("/files?client_id=" + client_id, token=outsider)[0], 404)
        self.assertEqual(self.call("/intake", "PATCH", {"enabled": 0}, owner)[0], 200)
        self.assertEqual(self.call("/public/intake?token=" + token)[0], 404)

    def test_unknown_intake_tokens_are_rate_limited(self):
        for index in range(60):
            self.assertEqual(
                self.call("/public/intake", "POST", {"token": f"missing-{index}"})[0],
                404,
            )
        self.assertEqual(
            self.call("/public/intake", "POST", {"token": "missing-last"})[0],
            429,
        )

    def test_full_client_quote_project_partial_payments_expense_document(self):
        token, _ = self.account("lifecycle")
        code, c = self.call(
            "/clients",
            "POST",
            {"name": "Client", "type": "company", "company": "Studio"},
            token,
        )
        self.assertEqual(code, 201)
        q = self.quote(token, client_id=c["item"]["id"])
        public_token = q["public_url"].split("quote=")[1]
        self.assertEqual(self.call("/public/quote?token=" + public_token)[0], 404)
        q = self.publish(token, q)
        code, public = self.call("/public/quote?token=" + public_token)
        self.assertEqual(code, 200)
        self.assertNotIn("cost_price", json.dumps(public))
        self.assertNotIn("internal_cost", json.dumps(public))
        self.assertNotIn("markup", json.dumps(public))
        self.assertEqual(public["quote"]["approval_state"], "viewed")
        code, _ = self.call(
            "/public/accept",
            "POST",
            {"token": public_token, "version": 1, "name": "Client"},
        )
        self.assertEqual(code, 200)
        code, project = self.call("/quotes/" + q["id"] + "/project", "POST", {}, token)
        self.assertEqual(code, 201)
        p = project["project"]
        self.assertEqual(
            self.call("/quotes/" + q["id"] + "/project", "POST", {}, token)[1][
                "project"
            ]["id"],
            p["id"],
        )
        receipt = {
            "project_id": p["id"],
            "amount_kopecks": 10000,
            "method": "bank_transfer",
        }
        self.assertEqual(
            self.call("/receipts", "POST", receipt, token, key="lifecycle-payment-key")[
                0
            ],
            201,
        )
        self.assertEqual(
            self.call("/receipts", "POST", receipt, token, key="lifecycle-payment-key")[
                0
            ],
            200,
        )
        self.assertEqual(
            self.call(
                "/receipts",
                "POST",
                {**receipt, "amount_kopecks": 30000},
                token,
                key="another-payment-key",
            )[0],
            409,
        )
        self.assertEqual(
            self.call(
                "/expenses",
                "POST",
                {"project_id": p["id"], "amount_kopecks": 2000, "category": "Tools"},
                token,
                key="lifecycle-expense-key",
            )[0],
            201,
        )
        overview = self.call("/overview", token=token)[1]["currencies"][0]
        self.assertEqual(
            (overview["paid"], overview["unpaid"], overview["cash_profit"]),
            (10000, 20000, 8000),
        )
        self.assertEqual(
            self.call(
                "/projects/" + p["id"],
                "PATCH",
                {"revision": 1, "status": "completed"},
                token,
            )[0],
            200,
        )
        code, document = self.call(
            "/documents", "POST", {"quote_id": q["id"], "kind": "act"}, token
        )
        self.assertEqual(code, 201, document)
        code, content = self.call(
            "/documents/" + document["document"]["id"] + "/pdf", token=token, raw=True
        )
        self.assertEqual(code, 200)
        self.assertTrue(content.startswith(b"%PDF-"))
        from pypdf import PdfReader

        self.assertIn("Client", PdfReader(io.BytesIO(content)).pages[0].extract_text())

    def test_version_immutability_stale_response_and_optimistic_lock(self):
        token, _ = self.account("versions")
        q = self.publish(token, self.quote(token))
        link = q["public_url"].split("quote=")[1]
        code, _ = self.call(
            "/quotes/" + q["id"],
            "PATCH",
            {"revision": q["revision"] - 1, "title": "stale"},
            token,
        )
        self.assertEqual(code, 409)
        code, updated = self.call(
            "/quotes/" + q["id"],
            "PATCH",
            {
                "revision": q["revision"],
                "title": "Revised",
                "items": [{"name": "New", "unit_price": 70000}],
            },
            token,
        )
        self.assertEqual(code, 200)
        old = self.call("/public/quote?token=" + link)[1]["quote"]
        self.assertEqual((old["title"], old["amount_kopecks"]), ("Project", 30000))
        for private_key in ("id", "client_id", "internal_cost", "revision", "custom_fields", "view_count"):
            self.assertNotIn(private_key, old)
        self.assertEqual(
            self.call("/public/accept", "POST", {"token": link, "version": 1})[0], 409
        )
        latest = self.publish(token, updated["quote"])
        self.assertEqual(latest["published_version"], 2)
        self.assertEqual(
            self.call("/public/accept", "POST", {"token": link, "version": 1})[0], 409
        )
        self.assertEqual(
            self.call("/public/accept", "POST", {"token": link, "version": 2})[0], 200
        )
        code, current = self.call("/quotes/" + q["id"], token=token)
        self.assertEqual(
            self.call(
                "/quotes/" + q["id"],
                "PATCH",
                {"revision": current["quote"]["revision"], "title": "cannot edit"},
                token,
            )[0],
            409,
        )

    def test_compare_versions_is_private_read_only_and_validates_selectors(self):
        owner, _ = self.account("compare-owner")
        other, _ = self.account("compare-other")
        original = self.publish(owner, self.quote(owner))
        path = "/quotes/" + original["id"]
        body = {"revision": original["revision"], "items": [{**row, "name": "Renamed", "unit_price": row["unit_price"] + 1000} for row in original["items"]], "terms": "New terms"}
        updated = self.call(path, "PATCH", body, owner)[1]["quote"]
        status, result = self.call(path + "/compare?from=1&to=current", token=owner)
        self.assertEqual(status, 200, result)
        self.assertEqual(result["revision"], updated["revision"])
        self.assertEqual(result["summary"]["changed"], len(original["items"]))
        self.assertEqual(result["fields"][0]["field"], "terms")
        self.assertEqual(self.call(path, token=owner)[1]["quote"], updated)
        self.assertEqual(self.call(path + "/compare?from=1", token=other)[0], 404)
        self.assertEqual(self.call(path + "/compare?from=1")[0], 401)
        for query in ("from=current", "from=0", "from=-1", "from=1&to=evil", "from=9999999999"):
            self.assertEqual(self.call(path + "/compare?" + query, token=owner)[0], 400)
        self.assertEqual(self.call(path + "/compare?from=99", token=owner)[0], 404)
        latest = self.publish(owner, updated)
        status, immutable = self.call(path + "/compare?from=1&to=2", token=owner)
        self.assertEqual(status, 200)
        self.assertEqual(result["items"], immutable["items"])
        public = self.call("/public/quote?token=" + latest["public_url"].split("quote=")[1])[1]["quote"]
        self.assertTrue(all("line_id" not in row and "cost_price" not in row for row in public["items"]))

    def test_tenant_isolation_viewer_rbac_and_revocation(self):
        owner, _ = self.account("team-owner")
        viewer, viewer_id = self.account("team-viewer")
        wid = self.call("/workspace", token=owner)[1]["workspace"]["id"]
        q = self.quote(owner)
        self.assertEqual(self.call("/quotes/" + q["id"], token=viewer)[0], 404)
        self.assertEqual(self.call("/workspace", token=viewer, workspace=wid)[0], 404)
        self.assertEqual(
            self.call(
                "/workspace/members",
                "POST",
                {"email": "team-viewer@test.invalid", "role": "viewer"},
                owner,
            )[0],
            200,
        )
        self.assertEqual(
            self.call("/quotes/" + q["id"], token=viewer, workspace=wid)[0], 200
        )
        self.assertEqual(
            self.call("/clients", "POST", {"name": "Denied"}, viewer, wid)[0], 403
        )
        self.assertEqual(self.call("/intake", "POST", {}, viewer, wid)[0], 403)
        self.assertEqual(self.call("/intake", "PATCH", {"enabled": 0}, viewer, wid)[0], 403)
        attachment = {"quote_id": q["id"], "name": "scope.txt", "content": base64.b64encode(b"Private scope").decode()}
        file_status, file_result = self.call("/files", "POST", attachment, owner)
        self.assertEqual(file_status, 201, file_result)
        file_id = file_result["file"]["id"]
        self.assertEqual(self.call("/files/" + file_id, "DELETE", token=viewer, workspace=wid)[0], 403)
        self.assertEqual(self.call("/files/" + file_id, token=viewer, workspace=wid, raw=True)[0], 200)
        self.assertEqual(
            self.call(
                "/workspace/members",
                "POST",
                {"email": "team-viewer@test.invalid", "role": "admin"},
                viewer,
                wid,
            )[0],
            403,
        )
        self.assertEqual(
            self.call("/workspace/members/" + viewer_id, "DELETE", token=owner)[0], 200
        )
        self.assertEqual(
            self.call("/quotes/" + q["id"], token=viewer, workspace=wid)[0], 404
        )

    def test_parallel_creation_cannot_bypass_free_quota(self):
        token, _ = self.account("concurrent")
        payload = {"title": "Parallel", "client": "Client", "amount": 1000}
        with concurrent.futures.ThreadPoolExecutor(max_workers=12) as pool:
            codes = list(
                pool.map(
                    lambda _: self.call("/quotes", "POST", payload, token)[0], range(14)
                )
            )
        self.assertEqual(codes.count(201), 10)
        self.assertEqual(codes.count(402), 4)

    def test_csv_atomicity_and_spreadsheet_formula_escape(self):
        token, _ = self.account("import")
        code, _ = self.call(
            "/transfer/catalog",
            "POST",
            {"csv": "name,price\nGood,100\nBad,nope"},
            token,
        )
        self.assertEqual(code, 400)
        self.assertEqual(self.call("/catalog", token=token)[1]["items"], [])
        self.assertEqual(
            self.call(
                "/transfer/clients",
                "POST",
                {"csv": 'name,email\n=HYPERLINK("bad"),client@test.invalid'},
                token,
            )[0],
            201,
        )
        code, csv = self.call("/transfer/clients", token=token, raw=True)
        self.assertEqual(code, 200)
        self.assertIn(b"'=HYPERLINK", csv)

    def test_logs_redact_public_tokens(self):
        token, _ = self.account("logging")
        q = self.publish(token, self.quote(token))
        secret = q["public_url"].split("quote=")[1]
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            self.call("/public/quote?token=" + secret)
        self.assertNotIn(secret, output.getvalue())
        self.assertIn("request_id", output.getvalue())

    def test_private_attachment_public_opt_in_and_invalid_format(self):
        owner, _ = self.account("file-owner")
        other, _ = self.account("file-other")
        q = self.publish(owner, self.quote(owner))
        link = q["public_url"].split("quote=")[1]
        payload = {
            "quote_id": q["id"],
            "name": "notes.txt",
            "content": base64.b64encode(b"Customer notes").decode(),
        }
        code, private = self.call("/files", "POST", payload, owner)
        self.assertEqual(code, 201, private)
        fid = private["file"]["id"]
        self.assertEqual(self.call("/files/" + fid, token=other)[0], 404)
        self.assertEqual(self.call("/public/file?token=" + link + "&id=" + fid)[0], 404)
        self.assertEqual(
            self.call("/files", "POST", {**payload, "name": "bad.png"}, owner)[0], 400
        )
        code, shared = self.call("/files", "POST", {**payload, "public": 1}, owner)
        self.assertEqual(code, 201)
        public_id = shared["file"]["id"]
        code, content = self.call(
            "/public/file?token=" + link + "&id=" + public_id, raw=True
        )
        self.assertEqual((code, content), (200, b"Customer notes"))
        self.assertEqual(
            self.call("/files/" + public_id, "DELETE", token=owner)[0], 200
        )
        self.assertEqual(
            self.call("/public/file?token=" + link + "&id=" + public_id)[0], 404
        )

    def test_ai_disabled_is_explicit_and_does_not_create_quotes(self):
        token, _ = self.account("ai-disabled")
        saved_key = os.environ.pop("OPENROUTER_API_KEY", None)
        try:
            self.assertFalse(self.call("/capabilities", token=token)[1]["ai_drafting"])
            self.assertEqual(
                self.call("/ai/draft", "POST", {"text": "Create a project"}, token)[0],
                503,
            )
            self.assertEqual(self.call("/quotes", token=token)[1]["quotes"], [])
        finally:
            if saved_key is not None:
                os.environ["OPENROUTER_API_KEY"] = saved_key

    def test_custom_field_validation_and_cross_workspace_references(self):
        one, _ = self.account("custom-owner")
        two, _ = self.account("custom-other")
        foreign = self.call("/clients", "POST", {"name": "Other"}, two)[1]["item"]["id"]
        self.assertEqual(
            self.call(
                "/projects",
                "POST",
                {"name": "Invalid", "amount_kopecks": 100, "client_id": foreign},
                one,
            )[0],
            404,
        )
        code, field = self.call(
            "/custom-fields",
            "POST",
            {
                "name": "Level",
                "type": "select",
                "entity_type": "clients",
                "options": ["A", "B"],
            },
            one,
        )
        self.assertEqual(code, 201)
        self.assertEqual(
            self.call(
                "/clients",
                "POST",
                {"name": "Test", "custom_fields": {field["id"]: "C"}},
                one,
            )[0],
            400,
        )
        self.assertEqual(
            self.call(
                "/clients",
                "POST",
                {"name": "Test", "custom_fields": {field["id"]: "A"}},
                one,
            )[0],
            201,
        )


if __name__ == "__main__":
    unittest.main()
